"""Build this single report using the read-only skill and source log."""
import sys
sys.dont_write_bytecode = True
import hashlib
import html
import json
import os
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'artifacts'
os.chdir(ROOT)
os.environ['MPLCONFIGDIR'] = str(OUT / 'mpl')
os.environ['MPLBACKEND'] = 'Agg'
sys.path.insert(0, str(ROOT / 'skill' / 'scripts'))
sys.path.insert(0, str(ROOT / 'skill'))
from parse_intent import parse_intent
from data_source_adapter import load_records, quality_check
from analyze_core import analyze, scenario_records
from constants import SCENARIOS, SCENARIO_CN, SCENARIO_COLORS
from templates.single_report import build_single_report_data
from render_core import render_md


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def inventory():
    return {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for directory in ('skill', 'data') for p in (ROOT / directory).rglob('*') if p.is_file()}


baseline = inventory()
for name in ('data', 'docs', 'charts'):
    (OUT / name).mkdir(exist_ok=True)
write_json(OUT / 'data' / 'inputs.json', baseline)
source = ROOT / 'data' / 'single.log'
raw = source.read_text(encoding='utf-8')
intent = parse_intent('读取 data/single.log，生成 single 性能报告，保留源数据中的产品、测试时长与全部性能指标。')
assert intent['data_source_type'] == 'local_file' and intent['report_type'] == 'single'
intent['other_info'] = '生成时间锚点：2026-06-19T16:00:00+08:00；保留全部原始性能指标。'
write_json(OUT / 'data' / 'intent.json', intent)

# The generic parser handles metrics; anchored header matches preserve this log's metadata.
extracted = load_records(intent)
product, endpoint = re.search(r'^DB:\s*(\S+)\s+host=(\S+)\s*$', raw, re.M).groups()
started = re.search(r'^Time:\s*(.+)$', raw, re.M).group(1).strip()
duration = int(re.search(r'^Duration:\s*(\d+)s\s*$', raw, re.M).group(1))
tables, table_size, prepared = re.search(r'^Data:\s*tables=(\d+),\s*table_size=(\d+),\s*prepare=(\w+)', raw, re.M).groups()
threads = list(map(int, re.search(r'^Threads:\s*([\d ]+)$', raw, re.M).group(1).split()))
record_pattern = re.compile(r'^---\s+(oltp_\w+)\s*\|\s*threads=(\d+)\s*\|\s*([\d:]+)\s*\.\.\.\s*TPS=([\d.]+),\s*QPS=([\d.]+),\s*P95=([\d.]+)ms,\s*P99=([\d.]+)ms\s*$', re.M)
source_rows = list(record_pattern.finditer(raw))
assert len(source_rows) == len(extracted['records']) > 0
source_index = {}
for m in source_rows:
    s, t, clock, *values = m.groups()
    key = (s, int(t))
    assert key not in source_index
    source_index[key] = {'values': list(map(float, values)), 'line': raw[:m.start()].count('\n') + 1, 'time': clock}
for record in extracted['records']:
    item = source_index[(record['scenario'], record['threads'])]
    assert [record[k] for k in ('tps', 'qps', 'p95_ms', 'p99_ms')] == item['values']
    record.update(product=product, duration_sec=duration, sample_time=item['time'], source_line=item['line'])
meta = extracted['meta']
meta['products'] = [product]
meta['scenarios'] = [s for s in SCENARIOS if s in meta['scenarios']]
meta['log_meta'] = [{
    'product': product, 'log_file': 'single.log', 'engine': 'sysbench',
    'start_time': started, 'db_endpoint': endpoint,
    'data_config': f'{tables} tables × {table_size} rows',
    'tables': int(tables), 'table_size': int(table_size), 'prepare': prepared,
    'duration_sec': duration, 'raw_header': raw.split('\n\n')[0],
}]
meta['source_info'].update(value='data/single.log', rows_fetched=len(source_rows), sha256=baseline['data/single.log'])
qc = quality_check(extracted)
meta['concurrencies_by_product_scenario'] = qc['concurrencies_by_product_scenario']
assert meta['concurrencies'] == threads
assert len(extracted['records']) == len(meta['products']) * len(meta['scenarios']) * len(threads)
assert all(coverage == threads for coverage in qc['concurrencies_by_product_scenario'][product].values())
write_json(OUT / 'data' / 'extracted_data.json', extracted)
write_json(OUT / 'data' / 'quality.json', qc)

analysis, _ = analyze(extracted)
analysis['report_type'] = 'single'
analysis['fairness_warnings'] = [
    {'level': 'INFO', 'text': '单产品、单数据集、各场景并发档完整，日志统一标注 Duration。'},
    {'level': 'WEAK', 'text': '源日志未提供 CPU、内存、磁盘、操作系统及数据库参数，硬件环境无法核查。'},
]
sp = analysis['single_product'][product]
sp['scale'] = {}
sp['latency'] = {}
sp['scenario_stats'] = {}
insight_rows = []
for s in meta['scenarios']:
    rows = scenario_records(extracted['records'], product, s)
    first, last = rows[0], rows[-1]
    peak = max(rows, key=lambda r: r['qps'])
    prefix = f'analysis_results.single_product.{product}.scenario_stats.{s}'
    stats = {
        'min_threads': first['threads'], 'max_threads': last['threads'],
        'first_qps': first['qps'], 'last_qps': last['qps'],
        'peak_qps': peak['qps'], 'peak_threads': peak['threads'],
        'peak_p95_ms': peak['p95_ms'], 'peak_p99_ms': peak['p99_ms'],
        'peak_to_first_ratio': peak['qps'] / first['qps'],
        'last_to_first_ratio': last['qps'] / first['qps'],
        'last_vs_peak_pct': (last['qps'] / peak['qps'] - 1) * 100,
        'p95_first': first['p95_ms'], 'p95_last': last['p95_ms'],
        'p99_first': first['p99_ms'], 'p99_last': last['p99_ms'],
        'p95_last_vs_peak_ratio': last['p95_ms'] / peak['p95_ms'],
    }
    sp['scenario_stats'][s] = stats
    sp['scale'][s] = [{'threads': r['threads'], 'qps': r['qps'], 'tps': r['tps']} for r in rows]
    sp['latency'][s] = [{'threads': r['threads'], 'p95_ms': r['p95_ms'], 'p99_ms': r['p99_ms']} for r in rows]
    for summary in sp['peak_summary']:
        if summary['scenario'] == s:
            summary['peak_p99_ms'] = peak['p99_ms']
    insight_rows.append({
        'level': 'L1', 'scenario': s, 'source': prefix,
        'text': f'{SCENARIO_CN[s]}（{s}）：{first["threads"]}→{last["threads"]:,} 并发，QPS 从 {first["qps"]:,.2f} 升至峰值 {peak["qps"]:,.2f}（峰值/起点 {stats["peak_to_first_ratio"]:.2f} 倍，{peak["threads"]} 并发）；最高并发 QPS 为 {last["qps"]:,.2f}，相对峰值 {stats["last_vs_peak_pct"]:+.2f}%。',
    })
    if peak['threads'] < last['threads'] and last['qps'] / peak['qps'] < 0.95:
        insight_rows.append({
            'level': 'L2', 'scenario': s, 'source': prefix,
            'text': f'{SCENARIO_CN[s]}在 {peak["threads"]} 并发后出现已测档位的性能拐点：{last["threads"]:,} 并发 QPS 较峰值下降 {-stats["last_vs_peak_pct"]:.2f}%（超过 5% 判据）。该发现不代表已定位瓶颈原因。',
        })
insights = {'report_type': 'single', 'single_product_insights': {product: insight_rows}, 'selection_insights': []}
assert all(i['level'] in ('L1', 'L2') and i['source'] for i in insight_rows)
write_json(OUT / 'data' / 'analysis_results.json', analysis)
write_json(OUT / 'data' / 'insights.json', insights)

# Native single charts, with an actual text-box intersection check before saving.
import generate_charts as charts
import matplotlib.pyplot as plt
from matplotlib.text import Text
plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'DejaVu Sans']
chart_checks = []
savefig = plt.savefig


def check_text_overlap(fig):
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    boxes = []
    for artist in fig.findobj(Text):
        if artist.get_visible() and artist.get_text().strip():
            box = artist.get_window_extent(renderer)
            if box.width > 0 and box.height > 0:
                boxes.append((artist.get_text(), box))
    overlaps = []
    for i, (text_a, box_a) in enumerate(boxes):
        for text_b, box_b in boxes[i + 1:]:
            width = min(box_a.x1, box_b.x1) - max(box_a.x0, box_b.x0)
            height = min(box_a.y1, box_b.y1) - max(box_a.y0, box_b.y0)
            if width > 1 and height > 1:
                overlaps.append([text_a, text_b])
    return overlaps


def checked_save(path, *args, **kwargs):
    fig = plt.gcf()
    overlaps = check_text_overlap(fig)
    title = fig.axes[0].get_title()
    assert re.search(r'[\u4e00-\u9fff]', title), title
    assert not overlaps, (str(path), overlaps)
    chart_checks.append({'file': Path(path).name, 'title': title, 'text_overlaps': overlaps})
    savefig(path, *args, **kwargs)


plt.savefig = checked_save
charts.gen_single(extracted, analysis, str(OUT / 'charts'))
fig, ax = plt.subplots(figsize=(9, 5.5))
for s in meta['scenarios']:
    rows = scenario_records(extracted['records'], product, s)
    ax.plot([r['threads'] for r in rows], [r['qps'] for r in rows], marker='o',
            label=SCENARIO_CN[s], color=SCENARIO_COLORS[s])
ax.set_xscale('log', base=2)
ax.set_xticks(threads, [str(t) for t in threads])
ax.set_title('单产品各场景并发扩展曲线', fontsize=14)
ax.set_xlabel('并发数（threads）')
ax.set_ylabel('QPS')
ax.grid(True, alpha=.3, linestyle='--')
ax.legend(loc='upper left', bbox_to_anchor=(1.02, 1), borderaxespad=0)
fig.tight_layout()
plt.savefig(OUT / 'charts' / 'single_scale.png')
plt.close(fig)
write_json(OUT / 'data' / 'chart_checks.json', chart_checks)


def paragraph(text):
    return {'type': 'paragraph', 'text': text}


def table(headers, rows):
    return {'type': 'table', 'headers': headers, 'rows': rows}


def image(name, caption):
    return {'type': 'image', 'path': f'charts/{name}.png', 'caption': caption}


def insight_block(item):
    return {'type': 'insight', **item}


report = build_single_report_data(extracted, analysis, insights, intent)
report['cover'].update(title='single 单次性能测试报告', date='2026-06-19')
sections = report['sections']
sections[0]['subsections'][0]['blocks'] = [paragraph(
    f'{product} 在 {started} 开始的 sysbench OLTP 测试中，覆盖 {len(meta["scenarios"])} 个独立场景及 {len(threads)} 个并发档（{", ".join(map(str, threads))}），共 {len(source_rows)} 条记录。'
    f'原日志标注 Duration: {duration}s；本报告将其保留为单项测试时长，不将其当作整份日志总耗时。所有场景的已测峰值均出现在 {sp["peak_summary"][0]["peak_threads"]} 并发。各场景独立分析，不合并均值。')]
sections[0]['subsections'][1]['blocks'] = [table(
    ['场景', '峰值并发', '峰值 TPS', '峰值 QPS', 'P95@峰值 (ms)', 'P99@峰值 (ms)'],
    [[SCENARIO_CN[p['scenario']], f'{p["peak_threads"]:,}', f'{p["peak_tps"]:,.2f}',
      f'{p["peak_qps"]:,.2f}', f'{p["peak_p95_ms"]:,.2f}', f'{p["peak_p99_ms"]:,.2f}'] for p in sp['peak_summary']]
)]
sections[0]['subsections'][2]['blocks'] = [
    image('single_peak_summary', f'{product} 各独立场景的峰值 QPS；柱上标签四舍五入，精确数值见表'),
    *[insight_block(i) for i in insight_rows if i['level'] == 'L1'],
]
sections[1]['subsections'] = [
    {'id': 'ch2_1', 'title': '2.1 数据库配置', 'blocks': [table(['项', '源日志值'], [
        ['产品及版本', product], ['测试起始时间', started], ['DB 端点', endpoint],
        ['数据集', f'tables={tables}, table_size={table_size}, prepare={prepared}'],
        ['测试时长', f'{duration}s（Duration: {duration}s）'],
        ['报告生成时间', '2026-06-19T16:00:00+08:00（任务固定时间）'],
    ])]},
    {'id': 'ch2_2', 'title': '2.2 硬件环境', 'blocks': [paragraph('源日志未提供 CPU、内存、存储、网络、操作系统以及数据库运行参数；上述信息均记为未提供，不能据此判断硬件容量或资源瓶颈。')]},
    {'id': 'ch2_3', 'title': '2.3 测试方法', 'blocks': [
        paragraph(f'工具为 sysbench，工具版本未提供。并发档完整保留为 {threads}，日志统一标注测试时长 {duration}s。TPS 为事务每秒，QPS 为查询每秒，P95/P99 分别为第 95/99 百分位延迟（ms）。'),
        table(['场景字段', '说明'], [[s, SCENARIO_CN[s]] for s in meta['scenarios']]),
        paragraph('各场景具有不同事务组成，不能据场景间 QPS 大小直接判定业务优劣。源文件是汇总日志，没有提供完整 sysbench 命令、预热规则和重复测试轮次。'),
    ]},
    {'id': 'ch2_4', 'title': '2.4 数据质量与公平性', 'blocks': [
        table(['检查项', '结果'], [
            ['记录数与覆盖', f'{len(meta["products"])} 产品 × {len(meta["scenarios"])} 场景 × {len(threads)} 并发档 = {len(source_rows)} 条；无重复、无缺档'],
            ['TPS / QPS / P95 / P99 空值率', '0% / 0% / 0% / 0%'],
            ['原始指标核对', f'{len(source_rows)} 条 × 4 项指标逐项核对一致，保留两位小数'],
            ['数据来源', 'data/single.log'],
            ['数据集与时长', '单一日志头，全场景使用同一数据集和 Duration 标注'],
            ['公平性范围', '单产品报告；无跨产品公平性结论；硬件环境未提供（WEAK）'],
        ]),
        paragraph('测试时区未在日志注明，保留原始时间，不进行时区换算。记录中的时刻为源日志逐项标记，不推算实际结束时间或整轮总耗时。'),
    ]},
]
sections[2]['blocks'] = [
    paragraph(f'产品：{product}。每条记录的测试时长均保留为日志头标注的 {duration}s。以下为全部 {len(source_rows)} 条记录；TPS/QPS 保留两位小数与千分位，P95/P99 单位为 ms。'),
    table(['场景', '并发', '原始时刻', 'TPS', 'QPS', 'P95 (ms)', 'P99 (ms)'], [
        [r['scenario'], f'{r["threads"]:,}', r['sample_time'], *[f'{r[k]:,.2f}' for k in ('tps', 'qps', 'p95_ms', 'p99_ms')]]
        for s in meta['scenarios'] for r in scenario_records(extracted['records'], product, s)
    ]),
]
sections[3]['blocks'] = [image('single_scale', f'{product} 各场景独立 QPS 曲线（不作均值或综合评分）')]
sections[3]['subsections'] = []
for index, s in enumerate(meta['scenarios'], 1):
    st = sp['scenario_stats'][s]
    blocks = [
        image(f'single_qps_{s}', f'{SCENARIO_CN[s]} 并发-QPS 曲线；横轴为以 2 为底的对数刻度，覆盖全部并发档'),
        image(f'single_p95_{s}', f'{SCENARIO_CN[s]} 并发-P95 曲线；横轴为以 2 为底的对数刻度，覆盖全部并发档'),
        *[insight_block(i) for i in insight_rows if i['scenario'] == s],
        paragraph(f'延迟稳定性：从 {st["min_threads"]} 到 {st["max_threads"]:,} 并发，P95 {st["p95_first"]:,.2f}→{st["p95_last"]:,.2f} ms，P99 {st["p99_first"]:,.2f}→{st["p99_last"]:,.2f} ms；最高并发的 P95 为峰值 QPS 档位的 {st["p95_last_vs_peak_ratio"]:.2f} 倍。未提供延迟 SLA，故不判定达标与否。'),
    ]
    sections[3]['subsections'].append({'id': f'ch4_{index}', 'title': f'4.{index} {SCENARIO_CN[s]}（{s}）', 'blocks': blocks})
sections[4]['subsections'] = [
    {'id': 'ch5_1', 'title': '5.1 后续验证与调优建议', 'blocks': [
        paragraph(f'全部场景的已测 QPS 峰值位于 {sp["peak_summary"][0]["peak_threads"]} 并发。建议在该档位附近补充测试并发点，并结合业务延迟目标选择并发上限；本次数据不足以直接给出生产连接池上限。'),
        paragraph('建议在保持现有数据集与时长一致的条件下重复测试，并同步记录 CPU、内存、磁盘、锁等待及缓存命中数据，再判断吞吐回落原因。没有资源观测数据，不给出具体 Buffer Pool 比例或未经验证的数据库参数。'),
    ]},
    {'id': 'ch5_2', 'title': '5.2 测试边界', 'blocks': [
        paragraph(f'仅覆盖本日志的 {product}、单一数据集、{len(meta["scenarios"])} 个场景和 {len(threads)} 个并发档。未提供重复测试、误差范围、资源监控或线上业务验证数据；结论仅描述已测档位。'),
        paragraph('本报告不补造缺失数据，不由尾延迟变化推断具体瓶颈，也不把单次基准结果外推为生产容量或持续稳定性保证。'),
    ]},
]
report['sections'].append({'id': 'ch6', 'title': '第 6 章 附录', 'blocks': [], 'subsections': [
    {'id': 'ch6_a', 'title': '6.1 附录 A：sysbench 命令与原始头部', 'blocks': [
        paragraph('完整 sysbench 命令未提供。下表逐行保留原始日志头，不据汇总信息拼造执行命令。'),
        table(['原始头部行'], [[html.escape(line)] for line in meta['log_meta'][0]['raw_header'].splitlines()]),
    ]},
    {'id': 'ch6_b', 'title': '6.2 附录 B：来源与复核方法', 'blocks': [
        table(['项', '值'], [['原始文件', 'data/single.log'], ['SHA-256', baseline['data/single.log']],
            ['标准化数据', 'data/extracted_data.json（相对报告目录）'], ['分析与洞察', 'data/analysis_results.json；data/insights.json'],
            ['图表', f'charts/，共 {len(chart_checks)} 张'], ['逐条来源', '标准化 records 中的 source_line 和 sample_time 对应原日志行号与时刻']]),
        paragraph('计算方法：各场景独立取 max(QPS) 为已测峰值；扩展倍数 = 峰值 QPS / 最低并发 QPS；最高并发相对峰值变化 = (最高并发 QPS / 峰值 QPS − 1) × 100%。只有峰值出现在最高并发之前且最高并发 QPS 低于峰值的 95% 时标注 L2 拐点。'),
        paragraph('L1 表示 Python 计算的客观事实，L2 表示统计发现。报告未加入 L3 原因推测。所有原始性能数据均来自原日志，补充环境信息明确注明未提供。'),
    ]},
]})
write_json(OUT / 'data' / 'report_data.json', report)


def all_blocks():
    for sec in report['sections']:
        yield from sec.get('blocks', [])
        for sub in sec.get('subsections', []):
            yield from sub.get('blocks', [])


# Markdown first, using the skill renderer; fix levels and add traceable insight labels.
md = render_md(report)
md = re.sub(r'^## (第 \d+ 章 .+)$', r'# \1', md, flags=re.M)
md = re.sub(r'^### (\d+\.\d+ .+)$', r'## \1', md, flags=re.M)
for b in all_blocks():
    if b['type'] == 'insight':
        icon = '✅' if b['level'] == 'L1' else '📊'
        md = md.replace(f'> **[{b["level"]}]** {b["text"]}', f'> {icon} {b["level"]} | 数据来源：{b["source"]}\n>\n> {b["text"]}')
    elif b['type'] == 'image':
        md = md.replace(f'*{b["caption"]}*', f'*图：{b["caption"]}*')
(OUT / 'report.md').write_text(md, encoding='utf-8')

# Independent Word rendering with explicit font and table styles.
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
doc = Document()
for style in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[style].font.name = '微软雅黑'
    doc.styles[style].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
doc.styles['Normal'].font.size = Pt(10.5)


def add_run_with_font(p, text, bold=False, color=None):
    run = p.add_run(str(text))
    run.font.name = '微软雅黑'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def word_blocks(blocks):
    for b in blocks:
        kind = b['type']
        if kind == 'paragraph':
            add_run_with_font(doc.add_paragraph(), b['text'])
        elif kind == 'table':
            tab = doc.add_table(rows=1, cols=len(b['headers']))
            tab.style = 'Table Grid'
            for c, label in zip(tab.rows[0].cells, b['headers']):
                add_run_with_font(c.paragraphs[0], label, True, '1F4E79')
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'D9E2F3')
                c._tc.get_or_add_tcPr().append(shade)
            repeat = OxmlElement('w:tblHeader')
            tab.rows[0]._tr.get_or_add_trPr().append(repeat)
            for ri, row in enumerate(b['rows']):
                for cell, value in zip(tab.add_row().cells, row):
                    add_run_with_font(cell.paragraphs[0], html.unescape(str(value)))
                    if ri % 2 == 1:
                        shade = OxmlElement('w:shd')
                        shade.set(qn('w:fill'), 'F8FAFC')
                        cell._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph()
        elif kind == 'image':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run_with_font(p, '').add_picture(str(OUT / b['path']), width=Cm(15.5))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run_with_font(cap, f'图：{b["caption"]}', color='666666').italic = True
        elif kind == 'insight':
            p = doc.add_paragraph()
            color = {'L1': '1971C2', 'L2': 'E67700', 'L3': 'C92A2A'}[b['level']]
            label = {'L1': '✅ L1 确定性结论', 'L2': '📊 L2 统计性发现', 'L3': '💡 L3 假设性洞察'}[b['level']]
            border = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            for key, val in {'val': 'single', 'sz': '18', 'space': '8', 'color': color}.items():
                left.set(qn(f'w:{key}'), val)
            border.append(left)
            p._p.get_or_add_pPr().append(border)
            add_run_with_font(p, f'{label} | 数据来源：{b["source"]}\n', True, color)
            add_run_with_font(p, b['text'])


add_run_with_font(doc.add_paragraph(style='Title'), report['cover']['title'])
add_run_with_font(doc.add_paragraph(style='Subtitle'), report['cover']['subtitle'])
add_run_with_font(doc.add_paragraph(), '日期：2026-06-19  版本：v1.0')
doc.add_page_break()
for sec in report['sections']:
    add_run_with_font(doc.add_paragraph(style='Heading 1'), sec['title'])
    word_blocks(sec.get('blocks', []))
    for sub in sec.get('subsections', []):
        add_run_with_font(doc.add_paragraph(style='Heading 2'), sub['title'])
        word_blocks(sub.get('blocks', []))
doc.save(OUT / 'report.docx')

# Independent HTML rendering from the same structured report data, never via Markdown.
def esc(value):
    return html.escape(str(value))


def html_blocks(blocks):
    result = []
    for b in blocks:
        if b['type'] == 'paragraph':
            result.append(f'<p>{esc(b["text"])}</p>')
        elif b['type'] == 'table':
            head = ''.join(f'<th>{esc(h)}</th>' for h in b['headers'])
            rows = ''.join('<tr>' + ''.join(f'<td>{esc(html.unescape(str(c)))}</td>' for c in row) + '</tr>' for row in b['rows'])
            result.append(f'<div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table></div>')
        elif b['type'] == 'image':
            result.append(f'<figure><img src="{esc(b["path"])}" alt="{esc(b["caption"])}"><figcaption>图：{esc(b["caption"])}</figcaption></figure>')
        elif b['type'] == 'insight':
            result.append(f'<aside class="insight {b["level"]}"><strong>{b["level"]} | 数据来源：{esc(b["source"])}</strong><p>{esc(b["text"])}</p></aside>')
    return '\n'.join(result)


css = """
body{font-family:'Microsoft YaHei','PingFang SC','Heiti TC','Hiragino Sans GB','STHeiti','Arial Unicode MS',sans-serif;max-width:1100px;margin:auto;padding:28px;color:#243447;line-height:1.7}
h1{color:#1f4e79;border-bottom:2px solid #1f4e79;padding-bottom:8px}h2{color:#2e75b6;border-left:4px solid #2e75b6;padding-left:12px}h3{color:#4a6d8c}
.table-wrap{overflow-x:auto}table{border-collapse:collapse;width:100%;margin:16px 0;font-size:14px}th{background:linear-gradient(#D9E2F3,#C5D3EE);color:#1f4e79}td,th{border:1px solid #ccd6e0;padding:8px;text-align:left}tbody tr:nth-child(even){background:#F8FAFC}
img{max-width:100%;height:auto}figure{margin:24px 0;text-align:center}figcaption{color:#666;font-style:italic}.insight{padding:12px 18px;margin:16px 0;overflow-wrap:anywhere}.L1{background:#E7F5FF;border-left:4px solid #1971C2}.L2{background:#FFF9DB;border-left:4px solid #F08C00}.L3{background:#FFE3E3;border-left:4px solid #C92A2A}nav{padding:16px;background:#F8FAFC}a{color:#1f4e79}
"""
toc = '<nav aria-label="目录"><h2>目录</h2><ul>' + ''.join(f'<li><a href="#{s["id"]}">{esc(s["title"])}</a></li>' for s in report['sections']) + '</ul></nav>'
body = [f'<header><h1>{esc(report["cover"]["title"])}</h1><p>{esc(report["cover"]["subtitle"])}</p><p>日期：2026-06-19 | 版本：v1.0</p></header>', toc]
for sec in report['sections']:
    body.append(f'<section><h1 id="{sec["id"]}">{esc(sec["title"])}</h1>')
    body.append(html_blocks(sec.get('blocks', [])))
    for sub in sec.get('subsections', []):
        body.append(f'<h2 id="{sub["id"]}">{esc(sub["title"])}</h2>')
        body.append(html_blocks(sub.get('blocks', [])))
    body.append('</section>')
html_text = '<!DOCTYPE html>\n<html lang="zh-CN"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>single 单次性能测试报告</title><style>' + css + '</style></head><body>' + '\n'.join(body) + '</body></html>'
(OUT / 'report.html').write_text(html_text, encoding='utf-8')

block_map = ['# 数据块映射', '', '| 原日志行 | 产品 | 场景 | 并发 | 原始时刻 | 报告位置 |', '| --- | --- | --- | --- | --- | --- |']
for r in extracted['records']:
    block_map.append(f'| {r["source_line"]} | {r["product"]} | {r["scenario"]} | {r["threads"]} | {r["sample_time"]} | 第3章全量数据；第4章对应场景 |')
(OUT / 'data' / 'block_map.md').write_text('\n'.join(block_map), encoding='utf-8')
(OUT / 'docs' / 'analysis_proposal.md').write_text(
    '# 分析方案与执行记录\n\n报告类型：single。输入：data/single.log。固定时间：2026-06-19T16:00:00+08:00。\n\n'
    '- ✅ 读取 SKILL.md 与全部六份 reference；自动识别数据源和报告意图。\n'
    '- ✅ 使用内置解析器提取指标，按源头部的 DB / Time / Duration 锚定元数据；技能源码未修改。\n'
    '- ✅ 全量核对 30 条记录及四项性能指标、产品、测试时长和并发档；保留 P99。\n'
    '- ✅ 分场景计算峰值、扩展倍数、最高并发回落及 P95/P99 变化；L2 拐点采用 5% 判据。\n'
    '- ✅ 生成 5 张 QPS、5 张 P95、1 张峰值汇总和 1 张单产品多场景曲线；检查文字框重叠。\n'
    '- ✅ 依次独立生成 Markdown、Word、HTML，共享结构化数据，不用格式转换工具。\n'
    '- ✅ 按最小交付标准进行文件、数据、图表、分析、排版与一致性核查，结果见 validation.md。\n', encoding='utf-8')


def validate():
    checks = []

    def check(name, condition, details):
        assert condition, f'{name}: {details}'
        checks.append((name, details))

    for name in ('intent', 'extracted_data', 'analysis_results', 'insights'):
        data = json.loads((OUT / 'data' / f'{name}.json').read_text(encoding='utf-8'))
        check(f'A JSON {name}', bool(data), '存在且 JSON 合法')
    for path in ('data/block_map.md', 'docs/analysis_proposal.md'):
        check(f'A {path}', (OUT / path).is_file(), '文件存在')
    for name, minimum in [('report.md', 5*1024), ('report.docx', 10*1024), ('report.html', 8*1024)]:
        size = (OUT / name).stat().st_size
        check(f'A {name}', size >= minimum, f'{size:,} 字节，达到基线')
    check('A Todo', '[ ]' not in (OUT / 'docs' / 'analysis_proposal.md').read_text(encoding='utf-8'), '执行清单全 ✅')
    check('B 产品、时长及记录覆盖', meta['products'] == [product] and meta['log_meta'][0]['duration_sec'] == duration and len(source_rows) == len(meta['scenarios'])*len(threads), f'{product}；{duration}s；{len(source_rows)} 条')
    check('B 指标空值率', all(qc[f'{k}_null_rate'] == 0 for k in ('tps', 'qps', 'p95', 'p99')), '四项指标空值率均为 0%')
    source_info = meta['source_info']
    check('B 来源与并发', all(k in source_info for k in ('type', 'value', 'rows_fetched')) and all(v == threads for v in qc['concurrencies_by_product_scenario'][product].values()), '来源字段完整，全部场景全并发覆盖')
    data_table = next(b for b in sections[2]['blocks'] if b['type'] == 'table')
    check('B 全量数据表', len(data_table['rows']) == len(source_rows), '第3章包含全部记录与 TPS/QPS/P95/P99')
    full_md = (OUT / 'report.md').read_text(encoding='utf-8')
    full_html = (OUT / 'report.html').read_text(encoding='utf-8')
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    with ZipFile(OUT / 'report.docx') as zipped:
        document_xml = ET.fromstring(zipped.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        word_text = '\n'.join(t.text or '' for t in document_xml.findall('.//w:t', ns))
        embedded = [f for f in zipped.namelist() if f.startswith('word/media/')]
    for r in extracted['records']:
        reference = source_index[(r['scenario'], r['threads'])]
        assert [r[k] for k in ('tps', 'qps', 'p95_ms', 'p99_ms')] == reference['values']
        for k in ('tps', 'qps', 'p95_ms', 'p99_ms'):
            value = f'{r[k]:,.2f}'
            assert all(value in text for text in (full_md, full_html, word_text)), (r, value)
    check('B/F 全量指标与三格式一致性', True, f'{len(source_rows)} 条 × 4 项指标逐项与原日志一致，且全部数值在三格式中存在；超过6条抽检要求')
    for text in (full_md, full_html, word_text):
        assert product in text and f'{duration}s' in text and '2026-06-19' in text
        assert not re.search(r'\{[A-Za-z_][A-Za-z0-9_]*\}', text)
    check('F 元数据与占位符', True, '产品、300s、固定报告日期在三格式一致；无未替换占位符')
    required = {f'single_{metric}_{s}.png' for metric in ('qps', 'p95') for s in meta['scenarios']}
    required |= {'single_peak_summary.png', 'single_scale.png'}
    actual = {p.name for p in (OUT / 'charts').glob('*.png')}
    check('A/C 图表完整性', required <= actual and len(chart_checks) == len(actual), f'{len(actual)} 张：全部场景 QPS/P95、峰值汇总和单产品曲线；排名图 N/A')
    check('C 中文标题与文字重叠', all(re.search(r'[\u4e00-\u9fff]', c['title']) and not c['text_overlaps'] for c in chart_checks), '12张均有中文标题；check_text_overlap 实际执行且全部返回空')
    image_paths = [b['path'] for b in all_blocks() if b['type'] == 'image']
    md_paths = re.findall(r'!\[[^\]]*\]\(([^)]+)\)', full_md)
    html_paths = re.findall(r'<img src="([^"]+)"', full_html)
    check('E 图片路径与Word嵌入', md_paths == image_paths and html_paths == image_paths and len(embedded) == len(image_paths) and all((OUT / p).exists() for p in image_paths), f'三格式均引用 {len(image_paths)} 张图，路径正确，Word嵌入完整')
    font_runs = document_xml.findall('.//w:r[w:t]', ns)
    check('E Word字体', len(font_runs) >= 5 and all(r.find('w:rPr/w:rFonts', ns).get(qn('w:eastAsia')) == '微软雅黑' for r in font_runs), f'全部 {len(font_runs)} 处文字 run 显式设置微软雅黑')
    word_tables = document_xml.findall('.//w:tbl', ns)
    check('E Word表头底色', len(word_tables) >= 3 and all(all(c.find('w:tcPr/w:shd', ns).get(qn('w:fill')) == 'D9E2F3' for c in t.findall('w:tr', ns)[0].findall('w:tc', ns)) for t in word_tables), f'全部 {len(word_tables)} 张表的表头为 D9E2F3')
    check('E HTML结构', all(s in full_html for s in ('<head>', 'charset=', 'viewport', '<style>', '<nav', 'Microsoft YaHei', 'PingFang SC', 'Heiti TC')), '完整 head、内联CSS、目录及中文字体回退')
    check('E Markdown层级与表格', len(re.findall(r'^# 第 \d+ 章', full_md, re.M)) == 6 and bool(re.search(r'^\| ---', full_md, re.M)), '章 #、节 ##，GFM表格含分隔行')
    core = sections[0]['subsections'][2]['blocks']
    check('D 核心结论', 3 <= sum(b['type'] == 'insight' for b in core) <= 5, '第1章共5条核心结论，每场景1条')
    check('D 洞察分级与来源', all(i['level'] in ('L1', 'L2') and i['source'] and i['scenario'] for i in insight_rows) and 'L3' not in json.dumps(insights), '所有洞察含level/source/scenario，无L3推测')
    check('D 场景分离与公平性', len(sections[3]['subsections']) == len(meta['scenarios']) and not any(w['level'] == 'STRONG' for w in analysis['fairness_warnings']), '5场景独立、无合并均值；环境缺失已披露，无STRONG未披露项')
    check('F 建议数值依据', '70%' not in full_md and '768' not in full_md, '移除默认参数建议；并发建议仅引用实测峰值并要求验证')
    check('输入与源码变更检查', baseline == inventory(), 'skill/ 与 data/ 文件清单及 SHA-256 均未变化；所有新增产物仅位于 artifacts/')
    (OUT / 'docs' / 'validation.md').write_text('# 最小交付核查\n\n固定时间：2026-06-19T16:00:00+08:00。\n\n' + '\n'.join(f'- ✅ {name}：{detail}' for name, detail in checks) + '\n\n图表文字检查记录：data/chart_checks.json。完整记录逐项对源核对，三格式数字核查通过。\n', encoding='utf-8')
    print(f'PASS: {len(checks)} validation groups; {len(source_rows)} records; {len(actual)} charts.')
    print(f'Product: {product}; Duration: {duration}s; report date: 2026-06-19.')
    print('Outputs: artifacts/report.md, artifacts/report.docx, artifacts/report.html')


validate()
