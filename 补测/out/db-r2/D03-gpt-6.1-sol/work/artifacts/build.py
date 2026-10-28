import sys
import os
os.environ['MPLCONFIGDIR'] = str(__import__('pathlib').Path.cwd() / 'artifacts' / 'mpl')
sys.dont_write_bytecode = True
from pathlib import Path
import json, hashlib, re, math
from itertools import product
ROOT = Path.cwd()
OUT = ROOT / 'artifacts'
for d in ('data','docs','charts'):
    (OUT/d).mkdir(exist_ok=True)
sys.path.insert(0, str(ROOT/'skill/scripts'))
sys.path.insert(0, str(ROOT/'skill'))
from detect_input import detect_data_source
from data_source_adapter import load_records, quality_check
from analyze_core import analyze
from analyze import analyze_iteration
from constants import SCENARIO_CN as CN, SCENARIO_COLORS as COLORS
from templates.iteration_report import build_iteration_report_data

def save(name, obj):
    (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()
before={str(p.relative_to(ROOT)):sha(p) for p in list((ROOT/'skill').rglob('*'))+[ROOT/'data/iter.json'] if p.is_file()}
source=detect_data_source('读取 data/iter.json 生成 iteration 历史趋势报告')
assert source['type']=='local_file'
raw=json.loads((ROOT/'data/iter.json').read_text(encoding='utf-8'))
versions=raw['meta']['products']
intent={'data_source_type':source['type'],'data_source_value':'data/iter.json','report_type':'iteration',
'dimension_filters':{'test_name_keywords':[],'tool_name':None,'tool_version_keywords':[], 'component_name':None,'component_version':None},
'results_key_filters':dict.fromkeys(['threads','tables','table_size','time','warehouses']),
 'task_filters':dict.fromkeys(['task_id','plan_id','report_id']),
'iteration_config':{'version_range':versions,'regression_threshold_pct':5},'custom_config':None,
'other_info':'按来源 meta.products 顺序计算相邻版本变化；下降严格超过 5% 为回退。固定当前时间 2026-06-19T16:00:00+08:00。'}
save('data/intent.json',intent)
e=load_records(intent)
qc=quality_check(e)
m=e['meta']; records=e['records']; scenarios=m['scenarios']; threads=m['concurrencies']
m['original_source_info']=m.get('source_info',{}).copy()
m['source_info']={'type':'local_file','value':'data/iter.json','rows_fetched':len(records),'sha256':sha(ROOT/'data/iter.json')}
m['concurrencies_by_product_scenario']=qc['concurrencies_by_product_scenario']
assert len(versions)>=3
idx={(r['product'],r['scenario'],r['threads']):r for r in records}
assert len(idx)==len(records)==len(versions)*len(scenarios)*len(threads)
assert set(idx)==set(product(versions,scenarios,threads))
assert all(math.isfinite(r[k]) and r[k]>0 for r in records for k in ('qps','tps','p95_ms','p99_ms'))
save('data/extracted_data.json',e);save('data/quality.json',qc)
a,_=analyze(e)
a['report_type']='iteration'; a['iteration']=analyze_iteration(e,intent)
it=a['iteration']
assert it['version_order']==versions
logs=m['log_meta']
assert len(set(x['data_config'] for x in logs))==1
assert len(set(x['duration_sec'] for x in logs))==1
assert m['test_env']
a['fairness_warnings']={'STRONG':[],'WEAK':['硬件仅提供共享 test_env，缺少逐版本环境快照；配置、缓存状态及重复运行信息未提供。'],'INFO':['数据集、测试时长、逐版本逐场景并发档一致。']}
it['p95_by_threads']={s:{str(t):[{'version':v,'p95_ms':idx[v,s,t]['p95_ms']} for v in versions] for t in threads} for s in scenarios}
for s in scenarios:
    for i,row in enumerate(it['trend'][s]):
        row['previous_version']=versions[i-1] if i else None
        # Independent recomputation from raw source records.
        peak=max(r['qps'] for r in raw['records'] if r['product']==versions[i] and r['scenario']==s)
        assert peak==row['peak_qps']
        if i:
            prev=max(r['qps'] for r in raw['records'] if r['product']==versions[i-1] and r['scenario']==s)
            assert math.isclose(row['delta_pct'],(peak/prev-1)*100,abs_tol=1e-10)
for r in it['regression_points']:
    r['previous_version']=versions[versions.index(r['version'])-1]
ins=[]
for s in scenarios:
    tr=it['trend'][s]
    ins.append({'level':'L1','scenario':s,'source':f'analysis_results.iteration.trend.{s}',
    'text':f"{CN[s]}（{s}）：{versions[0]} → {versions[-1]}，峰值 QPS {tr[0]['peak_qps']:,.2f} → {tr[-1]['peak_qps']:,.2f}，累计 {it['cumulative_change'][s]:+.2f}%。"})
for r in it['regression_points']:
    ins.append({'level':'L2','scenario':r['scenario'],'source':'analysis_results.iteration.regression_points',
    'text':f"回退：{r['previous_version']} → {r['version']}，{CN[r['scenario']]}峰值 QPS 下降 {abs(r['delta_pct']):.2f}%（严格超过 5%）。"})
insights={'single_product_insights':{},'iteration_insights':ins,'selection_insights':[]}
save('data/analysis_results.json',a);save('data/insights.json',insights)
assert all({'level','source','scenario','text'}<=x.keys() and x['level']!='L3' for x in ins)
# Charts use all six fixed concurrency levels for latency.
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
plt.rcParams.update({'font.sans-serif':['Microsoft YaHei','SimHei','DejaVu Sans'],'axes.unicode_minus':False,'font.size':10})
chart_checks={}
def check_text_overlap(fig):
    fig.canvas.draw(); renderer=fig.canvas.get_renderer()
    # Check visible titles, axis labels, ticks and legend entries, excluding empty text.
    texts=[]
    for ax in fig.axes:
        texts += [ax.title,ax.xaxis.label,ax.yaxis.label]+ax.get_xticklabels()+ax.get_yticklabels()
        if ax.get_legend(): texts += ax.get_legend().get_texts()
    boxes=[(t.get_text(),t.get_window_extent(renderer)) for t in texts if t.get_visible() and t.get_text().strip()]
    return [[a,c] for i,(a,b) in enumerate(boxes) for c,d in boxes[i+1:] if b.overlaps(d)]
def finish(fig, ax, name, title, ylabel, legend=False):
    ax.set_title(title,pad=18);ax.set_ylabel(ylabel,labelpad=12);ax.tick_params(axis='x',pad=12);ax.grid(alpha=.2)
    if legend: ax.legend(loc='upper left',bbox_to_anchor=(1.02,1),frameon=False)
    fig.tight_layout(pad=2)
    overlaps=check_text_overlap(fig)
    chart_checks[name]={'title':title,'check_text_overlap':overlaps}
    assert not overlaps,(name,overlaps)
    fig.savefig(OUT/'charts'/name,dpi=160)
    plt.close(fig)
labels=[v.split()[-1] for v in versions]
x=list(range(len(versions)))
for s in scenarios:
    fig,ax=plt.subplots(figsize=(10,5))
    ax.plot(x,[r['peak_qps'] for r in it['trend'][s]],'o-',color=COLORS[s]);ax.set_xticks(x,labels);ax.set_xlabel('版本')
    finish(fig,ax,f'iter_trend_qps_{s}.png',f'{CN[s]}峰值吞吐演进','峰值 QPS')
    fig,ax=plt.subplots(figsize=(11,5))
    for t in threads: ax.plot(x,[idx[v,s,t]['p95_ms'] for v in versions],'o-',label=f'{t} 并发')
    ax.set_xticks(x,labels);ax.set_xlabel('版本')
    finish(fig,ax,f'iter_trend_p95_{s}.png',f'{CN[s]}同并发延迟演进','P95（毫秒）',True)
fig,ax=plt.subplots(figsize=(10,5))
ax.bar([CN[s] for s in scenarios],[it['cumulative_change'][s] for s in scenarios],color=[COLORS[s] for s in scenarios]);ax.axhline(0,color='black',lw=.6)
finish(fig,ax,'iter_cumulative.png','首版至末版累计变化','峰值 QPS 变化率（%）')
fig,ax=plt.subplots(figsize=(11,5))
for j,s in enumerate(scenarios):
    xx=[i+(j-2)*.035 for i in x[1:]]; yy=[r['delta_pct'] for r in it['trend'][s][1:]]
    ax.scatter(xx,yy,label=CN[s],color=COLORS[s],s=45)
    for px,py in zip(xx,yy):
        if py < -5: ax.scatter([px],[py],s=130,facecolors='none',edgecolors='red')
ax.set_xlim(-0.3,3.3);ax.axhline(-5,color='red',linestyle='--',label='回退阈值：-5%');ax.set_xticks(x,labels);ax.set_xlabel('版本（红圈表示回退）')
finish(fig,ax,'iter_regression.png','相邻版本变化与回退点','峰值 QPS 变化率（%）',True)
fig,ax=plt.subplots(figsize=(11,5))
for s in scenarios:
    tr=it['trend'][s];ax.plot(x,[r['peak_qps']/tr[0]['peak_qps']*100 for r in tr],'o-',label=CN[s])
ax.set_xticks(x,labels);ax.set_xlabel('版本')
finish(fig,ax,'iter_summary.png','各场景归一化演进趋势','首版峰值为 100',True)
for i,v in enumerate(versions):
    fig,ax=plt.subplots(figsize=(11,5))
    for s in scenarios: ax.plot(threads,[idx[v,s,t]['qps'] for t in threads],'o-',label=CN[s])
    ax.set_xscale('log',base=2);ax.set_xticks(threads,[str(t) for t in threads]);ax.set_xlabel('并发数')
    finish(fig,ax,f'iter_version_{i+1}.png',f'{labels[i]} 各场景并发吞吐曲线','QPS',True)
save('data/chart_checks.json',chart_checks)
# Build on the supplied template; complete required sections in output data only.
rd=build_iteration_report_data(e,a,insights,intent)
rd['cover']['date']='2026-06-19'
P=lambda text:{'type':'paragraph','text':text}
T=lambda headers,rows:{'type':'table','headers':headers,'rows':rows}
I=lambda path,caption:{'type':'image','path':'charts/'+path,'caption':caption}
S=lambda sid,title,blocks:{'id':sid,'title':title,'blocks':blocks}
F=lambda n:f'{n:,.2f}'
sections=rd['sections']
sections[0]['subsections'][1]['blocks'] = [b for b in sections[0]['subsections'][1]['blocks'] if b['type'] != 'image']
sections[0]['subsections'][2]['blocks']=[dict(x,type='insight') for x in ins[:5]]+[dict(ins[[x['level'] for x in ins].index('L2')],type='insight')] if any(x['level']=='L2' for x in ins) else [dict(x,type='insight') for x in ins]
sections[1]['subsections']=[
 S('ch2_1','2.1 版本清单',[T(['版本','测试时间','任务 ID','数据集','时长（秒）'],[[r['product'],r['start_time'],r['task_id'],r['data_config'],str(r['duration_sec'])] for r in logs])]),
 S('ch2_2','2.2 硬件环境',[T(['项目','来源配置'],[[k,str(v)] for k,v in m['test_env'].items()]),P(a['fairness_warnings']['WEAK'][0])]),
 S('ch2_3','2.3 测试方法',[P('使用 sysbench 1.0.20。按来源 meta.products 列表顺序分析：'+' → '.join(versions)+'。各场景独立，不合并均值。扫描全部并发档 '+', '.join(map(str,threads))+' 的 QPS 最大值作为版本峰值。相邻变化率=(当前峰值−前版峰值)/前版峰值×100%；累计变化率=(末版峰值−首版峰值)/首版峰值×100%。仅未舍入变化率 < -5% 标记回退，恰好 -5% 不标记。P95 越低越好，按同场景同并发档分别比较，不用于吞吐回退判定。')]),
 S('ch2_4','2.4 数据质量与公平性',[P(f'共 {len(records)} 条记录 = {len(versions)} 版本 × {len(scenarios)} 场景 × {len(threads)} 并发档；键唯一且网格完整，TPS/QPS/P95/P99 空值率均为 0%。数据集均为 '+logs[0]['data_config']+f"，时长均为 {logs[0]['duration_sec']} 秒。公平性：并发档和场景逐项对齐，无 STRONG 警告；硬件一致性仅依据共享元数据，未经逐版本独立验证。实际读取 data/iter.json；输入中的旧 source_info 路径仅保留为来源声明，未访问。")])]
# Preserve exact source precision in full data tables.
for i,v in enumerate(versions):
    sections[2]['subsections'][i]['blocks'][0]['rows']=[[s,str(t),F(idx[v,s,t]['tps']),F(idx[v,s,t]['qps']),F(idx[v,s,t]['p95_ms']),F(idx[v,s,t]['p99_ms'])] for s in scenarios for t in threads]
    sections[2]['subsections'][i]['blocks'].append(I(f'iter_version_{i+1}.png',f'{v} 五场景曲线；每条线保留全部并发档，不计算跨场景均值'))
subs=[]
for i,s in enumerate(scenarios,1):
    tr=it['trend'][s]
    subs.append(S(f'ch4_{i}',f'4.{i} {CN[s]}（{s}）',[
      P(ins[i-1]['text']),T(['版本','峰值 QPS','峰值并发','相邻变化','标记'],[[r['version'],F(r['peak_qps']),str(r['peak_threads']),'-' if r['delta_pct'] is None else f"{r['delta_pct']:+.2f}%",'回退' if r['delta_pct'] is not None and r['delta_pct'] < -5 else '-'] for r in tr]),
      I(f'iter_trend_qps_{s}.png',f'{CN[s]}版本峰值趋势，版本顺序与来源一致'),
      T(['并发']+labels,[[str(t)]+[F(idx[v,s,t]['p95_ms']) for v in versions] for t in threads]),
      I(f'iter_trend_p95_{s}.png',f'{CN[s]} P95（毫秒）；六条线分别对应固定并发档，越低越好')]))
subs.append(S('ch4_cum','4.6 累计变化与综合趋势',[I('iter_cumulative.png','每根柱代表单独场景的首末版峰值变化'),I('iter_summary.png','每个场景分别以首版为 100 归一化，不形成综合评分')]))
subs.append(S('ch4_reg','4.7 回退点详细分析',[
 T(['前版本','当前版本','场景','前版峰值 QPS','当前峰值 QPS','相邻变化'],[[r['previous_version'],r['version'],r['scenario'],F(it['trend'][r['scenario']][versions.index(r['version'])-1]['peak_qps']),F(r['peak_qps']),f"{r['delta_pct']:+.2f}%"] for r in it['regression_points']]),I('iter_regression.png','全部相邻版本变化；红圈为严格低于 -5% 的回退点'),P('未提供版本变更说明或回退原因，因此不作原因归因。')]))
sections[3]['subsections']=subs
sections[4]['subsections']=[S('ch5_1','5.1 持续优化方向',[P('保留各场景和并发档的独立基线；在相同环境重复测试，确认累计变化可复现，并同步跟踪 P95/P99，避免只依据峰值选择版本。')]),S('ch5_2','5.2 回退修复建议',[P('优先复测峰值跌幅最大的回退场景，并在相同数据集、硬件、并发档和测试时长下核对前后版本变更。采集执行计划、资源使用和等待信息后再判断原因。已检测回退并不等同于已确认故障根因。')]),S('ch5_3','5.3 测试边界',[P('仅使用所给 120 条记录；无重复样本、误差区间、逐版本硬件快照、数据库详细参数和缓存状态，不能推断统计显著性或生产表现。回退只依据每场景版本峰值 QPS，不涵盖固定并发吞吐或延迟劣化；后两者请结合全量数据判断。报告时间以用户给定 2026-06-19T16:00:00+08:00 为准。')])]
sections.append({'id':'ch6','title':'第 6 章 附录','blocks':[],'subsections':[S('ch6_1','6.1 原始数据与可追溯性',[P('唯一实际输入：data/iter.json。SHA-256：'+m['source_info']['sha256']+'。标准化数据及来源旧声明见 data/extracted_data.json；精确计算结果见 data/analysis_results.json；质量检查见 data/quality.json。')]),S('ch6_2','6.2 测试命令与缺失信息',[P('输入仅提供 sysbench 版本、数据集和运行时长，未提供完整原始命令、连接参数、预热方式及运行次数，故不补写可执行命令。')])]})
save('data/report_data.json',rd)
(OUT/'docs/analysis_proposal.md').write_text('# 分析方案\n\n按 iteration 工作流，对 meta.products 的来源顺序执行每场景峰值、相邻变化、首末累计变化和严格下降超过 5% 的回退检测。P95 使用所有固定并发档。三格式共享结构数据并独立渲染，生成 17 张图。\n\n不进行原因推断；完整披露环境证据和重复样本缺失。\n',encoding='utf-8')
(OUT/'data/block_map.md').write_text('# 内容映射\n\n| 章节 | 数据来源 |\n| --- | --- |\n| 1 演进结论 | iteration.trend / cumulative_change / regression_points |\n| 2 环境 | meta.test_env / log_meta / quality.json |\n| 3 全量数据 | records，120 条 |\n| 4 场景分析 | iteration，固定并发 P95 |\n| 5 建议 | 已观察回退及缺失信息 |\n| 6 附录 | 实际来源与 SHA-256 |\n',encoding='utf-8')
print(json.dumps({'regressions':it['regression_points'],'cumulative':it['cumulative_change'],'charts':len(chart_checks)},ensure_ascii=False,indent=2))
# Rendering functions share report_data, never convert between formats.
import html
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ESC=lambda x:html.escape(str(x))
MARK={'L1':'✅','L2':'📊','L3':'💡'}
DESC={'L1':'确定性结论','L2':'统计性发现','L3':'假设性洞察'}
COLOR={'L1':'1971C2','L2':'E67700','L3':'C92A2A'}
def blocks():
    for sec in sections:
        yield 1,sec['id'],sec['title'],sec.get('blocks',[])
        for sub in sec['subsections']: yield 2,sub['id'],sub['title'],sub['blocks']
md=[rd['cover']['title'],'','> '+rd['cover']['subtitle'],'','日期：2026-06-19；报告版本：v1.0','']
for level,sid,title,bs in blocks():
    md+=['#'*level+' '+title,'']
    for b in bs:
        typ=b['type']
        if typ=='paragraph':md += [b['text'],'']
        elif typ=='table':
            md+=['| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |']
            md+=['| '+' | '.join(map(str,r))+' |' for r in b['rows']];md+=['']
        elif typ=='image': md += [f"![{b['caption']}]({b['path']})",f"*图：{b['caption']}*",'']
        elif typ=='insight':md += [f"> {MARK[b['level']]} {b['level']} | 数据来源：{b['source']}",'>',f"> {b['text']}",'']
(OUT/'report.md').write_text('\n'.join(md),encoding='utf-8')
doc=Document()
for sty in ('Normal','Title','Heading 1','Heading 2','Caption'):
    style=doc.styles[sty]; style.font.name='微软雅黑';style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
def add_run_with_font(p,text,bold=False,color=None):
    r=p.add_run(str(text));r.font.name='微软雅黑';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');r.bold=bold
    if color:r.font.color.rgb=RGBColor.from_string(color)
    return r
for text in (rd['cover']['title'],rd['cover']['subtitle'],'日期：2026-06-19；报告版本：v1.0'):
    add_run_with_font(doc.add_paragraph(),text)
for level,sid,title,bs in blocks():
    add_run_with_font(doc.add_paragraph(style=f'Heading {level}'),title)
    for b in bs:
        typ=b['type']
        if typ=='paragraph': add_run_with_font(doc.add_paragraph(),b['text'])
        elif typ=='table':
            tab=doc.add_table(rows=1,cols=len(b['headers']));tab.style='Table Grid'
            for i,h in enumerate(b['headers']):
                cell=tab.rows[0].cells[i];add_run_with_font(cell.paragraphs[0],h,True,'1F4E79')
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'D9E2F3');cell._tc.get_or_add_tcPr().append(sh)
            for ri,row in enumerate(b['rows'],1):
                cells=tab.add_row().cells
                for cell,val in zip(cells,row):
                    add_run_with_font(cell.paragraphs[0],val)
                    if ri%2==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F8FAFC');cell._tc.get_or_add_tcPr().append(sh)
        elif typ=='image':
            p=doc.add_paragraph();p.alignment=1;r=add_run_with_font(p,'');r.add_picture(str(OUT/b['path']),width=Cm(15.5))
            p=doc.add_paragraph();p.alignment=1;add_run_with_font(p,'图：'+b['caption'],color='666666').italic=True
        elif typ=='insight':
            p=doc.add_paragraph();add_run_with_font(p,f"{MARK[b['level']]} {b['level']} {DESC[b['level']]} | 数据来源：{b['source']}\n{b['text']}",color=COLOR[b['level']])
            pb=OxmlElement('w:pBdr');left=OxmlElement('w:left');left.set(qn('w:val'),'single');left.set(qn('w:sz'),'18');left.set(qn('w:color'),COLOR[b['level']]);pb.append(left);p._p.get_or_add_pPr().append(pb)
doc.save(OUT/'report.docx')
css="body{font-family:'Microsoft YaHei','PingFang SC','Heiti TC','Hiragino Sans GB','STHeiti','Arial Unicode MS',sans-serif;max-width:1100px;margin:auto;padding:24px;color:#333}h1{color:#1f4e79;border-bottom:2px solid}h2{color:#2e75b6;border-left:4px solid;padding-left:10px}h3{color:#4a6d8c}table{border-collapse:collapse;width:100%;margin:16px 0}td,th{border:1px solid #ccc;padding:7px}th{background:linear-gradient(#D9E2F3,#C5D3EE);color:#1F4E79}tr:nth-child(even){background:#F8FAFC}img{max-width:100%}.caption{text-align:center;color:#666;font-style:italic}.insight{padding:12px;margin:12px 0;border-left:4px solid}.L1{background:#E7F5FF;border-color:#1971C2}.L2{background:#FFF9DB;border-color:#F08C00}.L3{background:#FFE3E3;border-color:#C92A2A}"
ht=['<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+ESC(rd['cover']['title'])+'</title><style>'+css+'</style></head><body>', '<h1>'+ESC(rd['cover']['title'])+'</h1><p>'+ESC(rd['cover']['subtitle'])+'</p><p>日期：2026-06-19；报告版本：v1.0</p><nav class="toc"><h2>目录</h2><ul>']
ht += [f'<li><a href="#{sid}">{ESC(title)}</a></li>' for level,sid,title,bs in blocks()]
ht += ['</ul></nav>']
for level,sid,title,bs in blocks():
    ht += [f'<h{level} id="{sid}">{ESC(title)}</h{level}>']
    for b in bs:
        typ=b['type']
        if typ=='paragraph':ht += ['<p>'+ESC(b['text'])+'</p>']
        elif typ=='table':
            ht+=['<table><thead><tr>'+''.join('<th>'+ESC(h)+'</th>' for h in b['headers'])+'</tr></thead><tbody>']
            ht+=['<tr>'+''.join('<td>'+ESC(c)+'</td>' for c in r)+'</tr>' for r in b['rows']];ht+=['</tbody></table>']
        elif typ=='image': ht += [f'<img src="{b["path"]}" alt="{ESC(b["caption"])}"><p class="caption">图：{ESC(b["caption"])}</p>']
        elif typ=='insight':ht += [f'<div class="insight {b["level"]}">{MARK[b["level"]]} {b["level"]} {DESC[b["level"]]} | 数据来源：{ESC(b["source"])}<p>{ESC(b["text"])}</p></div>']
ht+=['</body></html>'];(OUT/'report.html').write_text('\n'.join(ht),encoding='utf-8')
# Delivery checks with recorded evidence.
checks=[]
def ck(label,condition,detail):
    checks.append({'item':label,'passed':bool(condition),'evidence':detail})
for name in ('intent','extracted_data','analysis_results','insights'):
    ck('A '+name,json.loads((OUT/f'data/{name}.json').read_text(encoding='utf-8')) is not None,'JSON 已重新解析')
for path in ('data/block_map.md','docs/analysis_proposal.md'):ck('A '+path,(OUT/path).is_file(),'文件存在')
for ext,minimum in [('md',5120),('docx',10240),('html',8192)]:ck('A report.'+ext,(OUT/f'report.{ext}').stat().st_size>=minimum,str((OUT/f'report.{ext}').stat().st_size)+' bytes')
ck('A6/C1-C6 charts',len(chart_checks)==17,'5 峰值 + 5 P95 + 累计 + 回退 + 归一化 + 4 版本曲线')
ck('C7/C8 chart text',all(not c['check_text_overlap'] for c in chart_checks.values()),'17 图标题中文；标题、轴标签、刻度、图例包围盒检测均无交叠')
ck('B1/B2/B5/B6',len(records)==120 and len(idx)==120 and qc['qps_null_rate']==qc['tps_null_rate']==qc['p95_null_rate']==0,'完整唯一 4×5×6 网格；逐项验证必需指标为正有限数')
ck('B3 full tables',all(len(sub['blocks'][0]['rows'])==30 for sub in sections[2]['subsections']),'4 张全量表各 30 条，保留原始两位小数')
ck('B4 source equality',records==raw['records'],'全部 120 条与原始记录逐字段一致，覆盖要求的 6 条抽检')
ck('B7 source',m['source_info']['value']=='data/iter.json','实际路径、记录数、SHA-256 已记录')
ck('D1-D3 insight',all(x['level'] in ['L1','L2'] and x['source'] for x in ins),'无 L3；第 1 章 6 条结论，5 条累计变化及 1 条回退')
ck('D4/D8 fairness',not a['fairness_warnings']['STRONG'],'数据集、时长一致；逐版本硬件证据不足已披露')
ck('D5/D7 scope',len(subs)==7,'五场景独立分析；六档并发全保留')
mdtext=(OUT/'report.md').read_text(encoding='utf-8');htmltext=(OUT/'report.html').read_text(encoding='utf-8')
refs=re.findall(r'!\[[^\]]*\]\(([^)]+)\)',mdtext)
ck('E1 Markdown',all((OUT/p).is_file() and p.startswith('charts/') for p in refs) and len(refs)==17,'章 #、节 ##；表含分隔行；17 个相对图片链接存在')
loaded=Document(OUT/'report.docx')
runs=[r for p in loaded.paragraphs for r in p.runs]+[r for t in loaded.tables for row in t.rows for cell in row.cells for p in cell.paragraphs for r in p.runs]
ck('E2.1 fonts',all(r.font.name=='微软雅黑' and r._element.rPr.rFonts.get(qn('w:eastAsia'))=='微软雅黑' for r in runs),'检查全部 '+str(len(runs))+' 个 run')
ck('E2.2 headers',all(c._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill'))=='D9E2F3' for t in loaded.tables for c in t.rows[0].cells),'全部 '+str(len(loaded.tables))+' 张表表头')
ck('E2.3 images',len(loaded.inline_shapes)==len(refs),'DOCX 图片数量 '+str(len(loaded.inline_shapes)))
ck('E3 HTML',all(x in htmltext for x in ['<head>','<style>','charset','viewport','class="toc"','PingFang SC','Microsoft YaHei','Heiti TC']) and re.findall(r'<img src="([^"]+)"',htmltext)==refs,'head/CSS/TOC/font fallback/17 图片路径')
ck('F1 placeholders',not re.search(r'\{[A-Za-z_][A-Za-z0-9_]*\}',mdtext+htmltext),'Markdown 与 HTML 占位符扫描为空')
doctext='\n'.join(p.text for p in loaded.paragraphs)+'\n'+'\n'.join(c.text for t in loaded.tables for row in t.rows for c in row.cells)
keynums=[F(r['peak_qps']) for s in scenarios for r in it['trend'][s]]
ck('F3 three formats',all(n in mdtext and n in htmltext and n in doctext for n in keynums),'20 个场景版本峰值跨三格式核对，按两位小数与千分位展示')
ck('source untouched',all(sha(ROOT/p)==digest for p,digest in before.items()),'skill 全部文件及原始 data/iter.json SHA-256 前后相同')
save('data/verification.json',{'checks':checks,'passed':all(c['passed'] for c in checks),'input_hashes':before})
lines=['# 最小交付核查','']+[f"- {'✅' if c['passed'] else '❌'} {c['item']}：{c['evidence']}" for c in checks]
(OUT/'docs/checklist.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
(OUT/'docs/todo.md').write_text('# Todo\n\n'+ '\n'.join('- ✅ '+s for s in ['意图与来源识别','接入与质量检查','分析与 CHECKPOINT③','图表生成与文字重叠核查','三格式独立渲染','A–F 核查及输入完整性检查'])+'\n',encoding='utf-8')
assert all(c['passed'] for c in checks),[c for c in checks if not c['passed']]
print('DELIVERY CHECKS PASSED:',len(checks))


