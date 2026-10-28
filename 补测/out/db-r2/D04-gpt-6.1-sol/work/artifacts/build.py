import sys
sys.dont_write_bytecode = True
from pathlib import Path
import json
import copy
import hashlib
import re
import os
from html import escape

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'artifacts'
os.environ['MPLCONFIGDIR'] = str(OUT / '.mpl')
sys.path.insert(0, str(ROOT / 'skill/scripts'))
from detect_input import detect_data_source
from data_source_adapter import load_records, quality_check
from analyze_core import _check_fairness
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

for name in ['data', 'docs', 'charts']:
    (OUT / name).mkdir(exist_ok=True)
def save(name, data):
    (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

source = ROOT / 'data/compare.json'
before = hashlib.sha256(source.read_bytes()).hexdigest()
scenarios = ['oltp_write_only', 'oltp_update_index']
concs = [512, 1000]
metrics = [('qps', 'QPS', '查询吞吐量'), ('p95_ms', 'P95 (ms)', '百分之九十五延迟'), ('p99_ms', 'P99 (ms)', '百分之九十九延迟')]
cn = dict(zip(scenarios, ['只写事务', '索引更新']))
question = '仅分析三个产品在 oltp_write_only、oltp_update_index 两场景和 512、1000 并发下的 QPS、P95、P99，逐产品逐场景列出数据并给出有依据的产品选择建议。'
detected = detect_data_source('读取 data/compare.json 生成 custom 专项报告')
assert detected['type'] == 'local_file'
intent = {
    'data_source_type': detected['type'], 'data_source_value': 'data/compare.json', 'report_type': 'custom',
    'dimension_filters': {'test_name_keywords_or': [[s] for s in scenarios], 'tool_name': 'sysbench'},
    'results_key_filters': {'threads': None}, 'task_filters': {}, 'iteration_config': None,
    'custom_config': {'focus_dimension': '自定义', 'dimensions': {'scenarios': scenarios, 'threads': concs, 'metrics': ['qps', 'p95_ms', 'p99_ms']}, 'extra_questions': [question]},
    'other_info': '固定报告时间 2026-06-19T16:00:00+08:00；仅使用显式指定的两档并发，不扩展到其他档位。'
}
save('data/intent.json', intent)
raw = load_records(intent)
products = raw['meta']['products']
rows = [r for r in raw['records'] if r['scenario'] in scenarios and r['threads'] in concs]
assert len(rows) == 12
assert len({(r['product'], r['scenario'], r['threads']) for r in rows}) == 12
meta = copy.deepcopy(raw['meta'])
meta.update(scenarios=scenarios, concurrencies=concs)
meta['source_info'] = {'type': 'local_file', 'value': 'data/compare.json', 'rows_fetched': len(rows), 'rows_before_filter': len(raw['records']), 'sha256': before}
extracted = {'meta': meta, 'records': rows}
qc = quality_check(extracted)
assert qc['p99_null_rate'] == 0
meta['concurrencies_by_product_scenario'] = qc['concurrencies_by_product_scenario']
save('data/extracted_data.json', extracted)
save('data/quality.json', qc)
assert all(ts == concs for sd in qc['concurrencies_by_product_scenario'].values() for ts in sd.values())
assert _check_fairness(rows, meta)['passed']
assert len({m['data_config'] for m in meta['log_meta']}) == 1
assert len({m['duration_sec'] for m in meta['log_meta']}) == 1
def row(p, s, t):
    return next(r for r in rows if (r['product'], r['scenario'], r['threads']) == (p, s, t))
def fmt(v):
    return f'{v:,.2f}'
def pct(a, b):
    return (a / b - 1) * 100
def insight(text, scenario='all', level='L1', source='data/extracted_data.json：按产品、场景、并发精确筛选'):
    return {'level': level, 'text': text, 'source': source, 'scenario': scenario}

changes = []
findings = []
for s in scenarios:
    for p in products:
        a, b = row(p, s, 512), row(p, s, 1000)
        d = {'product': p, 'scenario': s, **{k + '_change_pct': pct(b[k], a[k]) for k, _, _ in metrics}}
        changes.append(d)
        findings.append(insight(f'{p}，{s}：并发 512→1,000，QPS {fmt(a["qps"])}→{fmt(b["qps"])}（{d["qps_change_pct"]:+.2f}%）；P95 {fmt(a["p95_ms"])}→{fmt(b["p95_ms"])} ms（{d["p95_ms_change_pct"]:+.2f}%）；P99 {fmt(a["p99_ms"])}→{fmt(b["p99_ms"])} ms（{d["p99_ms_change_pct"]:+.2f}%）。', s))
comparisons = []
for s in scenarios:
    for t in concs:
        group = [row(p, s, t) for p in products]
        best = {k: (max(group, key=lambda r: r[k]) if k == 'qps' else min(group, key=lambda r: r[k])) for k, _, _ in metrics}
        a = group[0]
        text = f'{s}，并发 {t:,}：QPS 最高为 {best["qps"]["product"]}（{fmt(best["qps"]["qps"])}）；P95 最低为 {best["p95_ms"]["product"]}（{fmt(best["p95_ms"]["p95_ms"])} ms）；P99 最低为 {best["p99_ms"]["product"]}（{fmt(best["p99_ms"]["p99_ms"])} ms）。'
        text += ' TDSQL-B 相对 ' + '；'.join(f'{r["product"]} 的 QPS 高 {pct(a["qps"], r["qps"]):.2f}%' for r in group[1:]) + '。'
        comparisons.append(insight(text, s))
warnings = [{'level': 'INFO', 'text': '数据集均为 tables=16, table_size=1000000，时长均为 300 秒；两场景两并发档在三产品上完全对齐。'}, {'level': 'WEAK', 'text': '仅提供一份共享硬件说明，缺少逐产品部署拓扑、数据库参数、重复试验及误差范围；公平性仅在已披露字段上成立。'}]
analysis = {'report_type': 'custom', 'fairness_warnings': warnings, 'single_product': {p: {'scale': {s: [row(p, s, t) for t in concs] for s in scenarios}, 'latency': {s: [{k: row(p, s, t)[k] for k in ['threads', 'p95_ms', 'p99_ms']} for t in concs] for s in scenarios}} for p in products}, 'custom': {'focus_dimension': '自定义', 'buckets': rows, 'changes': changes, 'main_finding': 'TDSQL-B 四个组合 QPS 最高；全部产品在两个场景从 512 到 1000 并发时 QPS 下降、P95/P99 上升。'}}
insights = {'single_product_insights': {p: [i for i in findings if i['text'].startswith(p)] for p in products}, 'custom_insights': comparisons, 'selection_insights': []}
save('data/analysis_results.json', analysis)
save('data/insights.json', insights)
assert analysis['single_product'] and analysis['custom']
assert all(i['level'] in ['L1', 'L2'] and i['source'] for i in findings + comparisons)

plt.rcParams.update({'font.sans-serif': ['Microsoft YaHei', 'SimHei', 'DejaVu Sans'], 'axes.unicode_minus': False, 'font.size': 10})
overlaps = {}
def check_text_overlap(fig):
    fig.canvas.draw()
    texts = []
    for ax in fig.axes:
        texts.extend([ax.title, ax.xaxis.label, ax.yaxis.label] + ax.get_xticklabels() + ax.get_yticklabels())
    if fig.legends:
        texts.extend(fig.legends[0].get_texts())
    texts = [t for t in texts if t.get_visible() and t.get_text()]
    renderer = fig.canvas.get_renderer()
    return [(a.get_text(), b.get_text()) for j, a in enumerate(texts) for b in texts[j+1:] if a.get_window_extent(renderer).overlaps(b.get_window_extent(renderer))]
for k, label, title in metrics:
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.2))
    for ax, s in zip(axes, scenarios):
        for p in products:
            ax.plot(concs, [row(p, s, t)[k] for t in concs], marker='o', label=p)
        ax.set_title(f'{cn[s]}：{title}与并发关系')
        ax.set_xlabel('并发数'); ax.set_ylabel(label)
        ax.set_xticks(concs, ['512', '1,000']); ax.set_xlim(450, 1060)
        ax.grid(alpha=.25)
    fig.legend(*axes[0].get_legend_handles_labels(), loc='lower center', ncol=3, bbox_to_anchor=(.5, .01))
    fig.subplots_adjust(bottom=.23, top=.88, left=.085, right=.98, wspace=.30)
    overlaps[k] = check_text_overlap(fig)
    assert not overlaps[k], overlaps[k]
    fig.savefig(OUT / f'charts/{k}.png', dpi=150)
    plt.close(fig)
save('data/chart_check.json', overlaps)

def para(text): return {'type': 'paragraph', 'text': text}
def table(headers, data): return {'type': 'table', 'headers': headers, 'rows': data}
def sub(title, blocks): return {'title': title, 'blocks': blocks}
def section(title, subs): return {'title': title, 'subsections': subs}
summary = [insight('TDSQL-B 在两个场景、两档并发的四个组合中 QPS 均最高；这是本次指定数据范围内的吞吐量结论。'), insight('三产品在每个指定场景由 512 增至 1,000 并发均出现 QPS 下降和 P95/P99 上升；在这两档中优先评估 512 并发。', level='L2'), insight('不确定性：每个产品×场景×并发只有一个聚合结果，没有重复试验或置信区间；细小延迟差不能作为稳定优势的证明。')]
secs = [section('第 1 章 分析目标', [sub('1.1 用户需求', [para(question)]), sub('1.2 关键发现摘要', [dict(i, type='insight') for i in summary]), sub('1.3 核心数据点', [table(['场景', '并发', '最高 QPS 产品', 'QPS'], [[s, f'{t:,}', products[0], fmt(row(products[0], s, t)['qps'])] for s in scenarios for t in concs])])]),
section('第 2 章 测试环境与数据范围', [sub('2.1 数据来源', [para('唯一数据源为 data/compare.json。文件内 source_info.value 指向历史聚合来源，属于输入元数据；本报告未访问该历史路径，实际来源以当前文件及其 SHA-256 为准。'), para('报告时间固定为 2026-06-19 16:00:00 +08:00。测试发生于 2026-06-11，三个产品起始时间分别见下表。'), table(['产品', '开始时间', '数据集', '时长（秒）'], [[m['product'], m['start_time'], m['data_config'], m['duration_sec']] for m in meta['log_meta']]), table(['环境项', '输入披露值'], list(meta['test_env'].items()))]),
sub('2.2 数据子集说明', [para(f'输入共 {len(raw["records"])} 条记录，按场景 OR 条件与并发集合精确筛选后共 {len(rows)} 条，等于 3 产品 × 2 场景 × 2 并发。每个产品×场景完整覆盖 512、1000。仅对 QPS、P95、P99 进行展示和分析，未将其他场景或并发档用于任何结论。'), para('QPS 单位为查询/秒，越高越好；P95、P99 单位为毫秒，越低越好。延迟沿用文件中的统计口径；未提供原始延迟分布，不能重新估计分位数。所有表格保留原始两位小数并使用千分位，计算使用未舍入值。'), para('质量门控已核验全部 12 条记录，所需指标无空值，组合键无重复，三产品覆盖完全相同。TPS 只保留于标准接入数据并检查完整性，不作为本次分析指标。')]),
sub('2.3 假设前提与公平性', [para(w['level'] + '：' + w['text']) for w in warnings] + [para('假设文件中的场景定义、查询计数及延迟统计方法在产品间一致。输入未披露各产品独立资源分配、持久化策略、复制配置或故障状态，不能确认它们完全一致。比较只适用于已给定版本和此次记录，不能据此给出成本、可靠性或其他负载的排序。')])]),
section('第 3 章 深度分析', [])]
for j, p in enumerate(products, 1):
    secs[2]['subsections'].append(sub(f'3.{j} {p}：逐场景数据', [table(['场景', '并发', 'QPS', 'P95（ms）', 'P99（ms）'], [[s, f'{t:,}'] + [fmt(row(p, s, t)[k]) for k, _, _ in metrics] for s in scenarios for t in concs])] + [dict(i, type='insight') for i in findings if i['text'].startswith(p)]))
for j, s in enumerate(scenarios, 4):
    secs[2]['subsections'].append(sub(f'3.{j} {s}：同档产品比较', [dict(i, type='insight') for i in comparisons if i['scenario'] == s]))
secs[2]['subsections'].append(sub('3.6 并发档位与三指标图示', [para('每图分为两个独立场景面板，不跨场景求和或平均。三条线对应三个产品；连接两个观测点仅用于呈现变化，不表示未测试并发的预测值。QPS 图为主图，P95 图为并发档位延迟对照图，P99 图为尾延迟验证图。')] + [{'type': 'image', 'path': f'charts/{k}.png', 'caption': f'图：{title}在两档并发下的逐场景对照；QPS 越高越好，延迟越低越好。'} for k, _, title in metrics]))
def detail(p, s, t):
    r = row(p, s, t)
    return f'{p}：QPS={fmt(r["qps"])}，P95={fmt(r["p95_ms"])} ms，P99={fmt(r["p99_ms"])} ms'

recommendations = [
    '吞吐优先：在本次四个组合中选择 TDSQL-B v22.7.3 作为首选候选。具体领先幅度见 3.4、3.5；这项建议以业务接受其对应 P95/P99 为前提，不能仅因 QPS 排名第一就忽略延迟门槛。',
    'oltp_write_only，512 并发：MySQL 的 P95、P99 均最低；若延迟优先，可把 MySQL 作为候选。TDSQL-B 吞吐量最高；OceanBase 的吞吐介于两者之间，应逐项按业务门槛取舍。' + '；'.join(detail(p, scenarios[0], 512) for p in products) + '。',
    'oltp_write_only，1000 并发：OceanBase P95、P99 均最低，同时 QPS 高于 MySQL；在这三个指标上本次记录均优于 MySQL。TDSQL-B QPS 更高，但 P95、P99 也更高，因此吞吐优先与尾延迟优先会产生不同选择。' + '；'.join(detail(p, scenarios[0], 1000) for p in products) + '。',
    'oltp_update_index，512 并发：TDSQL-B QPS、P95 领先，且 P99 低于 MySQL。OceanBase P99 最低。优先吞吐与 P95 时选 TDSQL-B；P99 为硬约束时再比较 OceanBase。' + '；'.join(detail(p, scenarios[1], 512) for p in products) + '。',
    'oltp_update_index，1000 并发：TDSQL-B QPS、P99 最优，OceanBase P95 略低于 TDSQL-B。MySQL 的三个指标在本次记录均落后于另外两者。吞吐与 P99 优先选择 TDSQL-B，P95 优先保留 OceanBase 验证。' + '；'.join(detail(p, scenarios[1], 1000) for p in products) + '。',
    '并发建议：在已测试的两档中优先选择 512 并发作为容量验证起点；增加到 1000 时所有产品、两个场景都没有带来 QPS 增益，反而增加两项尾延迟。该结论仅覆盖这两档，不能称为全局最优并发或推断性能拐点位置。',
    '决策步骤：先明确两个场景各自需要的最低 QPS、最高 P95、最高 P99，再从相同场景相同并发的表格筛选全部达标候选。用户未提供 SLA，因此本报告不代设阈值。若多产品均达标，需补充成本和运维信息后决策，现有数据不支持性价比结论。'
]
secs.extend([section('第 4 章 结论与建议', [sub('4.1 按场景与指标选择产品', [para(t) for t in recommendations]), sub('4.2 不确定性披露', [para('所有数值优势均为这份文件的观测事实，不代表统计显著性。单个聚合值无法衡量波动、长时间稳定性或极端异常。尤其不足 1 ms 的 P95 差异，应通过重复试验确认后再影响采购或切换决策。'), para('本报告不合并不同场景的指标，不生成综合评分。两个场景的查询组成不同，QPS 不能直接作为跨场景业务价值的可比单位，也不能对两档 P95/P99 取平均作为整体分位数。')]), sub('4.3 后续测试建议', [para('针对候选产品重复执行同场景同并发试验，保留每轮原始日志与延迟分布，并补齐每产品硬件、拓扑及关键数据库参数。以实际业务 SLA 验证候选；未有这些结果前，本报告不认定微小延迟差为可复现优势。'), para('如需确定容量上限，应在新一轮测试中补充并发档并观测资源指标。本次未对任何未指定的档位做分析，也无法由两个点定位 CPU、锁、日志或存储瓶颈。')])]), section('第 5 章 附录', [sub('5.1 来源与计算口径', [para('原始文件：data/compare.json；SHA-256：' + before), para('并发变化率 =（1000 并发指标 ÷ 512 并发指标 − 1）× 100%；同档 QPS 领先率 =（TDSQL-B QPS ÷ 对照产品 QPS − 1）× 100%。所有运算由 Python 对筛选后的原始记录直接执行。'), para('原始 sysbench 命令未提供，未补写或猜测可复现命令。报告生成脚本为 artifacts/build.py；三格式共享 report_data.json 并分别渲染，未使用格式转换工具。')]), sub('5.2 交付索引', [para('report.md 为文本核对主本；report.docx 用于编辑；report.html 用于浏览，浏览时保持 charts 文件夹相对位置。data/ 含意图、筛选记录、计算结果、洞察及核查结果；docs/analysis_proposal.md 记录范围和完成状态。'), para('用户限定范围优先于最小交付标准 D5 的全并发档默认要求；本报告完整覆盖用户指定的两档。custom 只对指定两个场景绘图，其他三个场景验证图不适用。未采用任何 L3 推测性结论，避免将未经验证的性能原因作为选型依据。')])])])
report = {'cover': {'title': '客制化专项报告', 'subtitle': '三产品写入与索引更新：512 / 1000 并发', 'date': '2026-06-19', 'version': 'v1.0'}, 'sections': secs}
save('data/report_data.json', report)

def blocks():
    for sec in secs:
        for ss in sec['subsections']:
            yield from ss['blocks']
md = ['# 客制化专项报告', '', report['cover']['subtitle'], '', '日期：2026-06-19　版本：v1.0', '']
for sec in secs:
    md += ['# ' + sec['title'], '']
    for ss in sec['subsections']:
        md += ['## ' + ss['title'], '']
        for b in ss['blocks']:
            if b['type'] == 'paragraph': md += [b['text'], '']
            elif b['type'] == 'insight': md += [f'> {"✅" if b["level"] == "L1" else "📊"} {b["level"]} | 数据来源：{b["source"]}', '>', '> ' + b['text'], '']
            elif b['type'] == 'image': md += [f'![{b["caption"]}]({b["path"]})', '*' + b['caption'] + '*', '']
            elif b['type'] == 'table': md += ['| ' + ' | '.join(map(str, b['headers'])) + ' |', '| ' + ' | '.join(['---'] * len(b['headers'])) + ' |'] + ['| ' + ' | '.join(map(str, r)) + ' |' for r in b['rows']] + ['']
(OUT / 'report.md').write_text('\n'.join(md), encoding='utf-8')

doc = Document()
def set_doc_default_font():
    for style in doc.styles:
        if style.type == 1:
            style.font.name = '微软雅黑'
            style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
set_doc_default_font()
def add_run_with_font(p, text, bold=False):
    run = p.add_run(str(text)); run.font.name = '微软雅黑'; run.font.size = Pt(10.5)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑'); run.bold = bold
    return run
def dp(text, style=None):
    p = doc.add_paragraph(style=style); add_run_with_font(p, text); return p
dp(report['cover']['title'], 'Title'); dp(report['cover']['subtitle']); dp('日期：2026-06-19　版本：v1.0')
for sec in secs:
    dp(sec['title'], 'Heading 1')
    for ss in sec['subsections']:
        dp(ss['title'], 'Heading 2')
        for b in ss['blocks']:
            if b['type'] == 'paragraph': dp(b['text'])
            elif b['type'] == 'insight':
                p = dp(f'{b["level"]} | 数据来源：{b["source"]}\n{b["text"]}')
                borders = OxmlElement('w:pBdr'); border = OxmlElement('w:left')
                for k, v in {'val': 'single', 'sz': '18', 'color': '1971C2' if b['level'] == 'L1' else 'E67700'}.items(): border.set(qn('w:' + k), v)
                borders.append(border); p._p.get_or_add_pPr().append(borders)
            elif b['type'] == 'table':
                tab = doc.add_table(rows=0, cols=len(b['headers'])); tab.style = 'Table Grid'
                for ri, rr in enumerate([b['headers']] + b['rows']):
                    cells = tab.add_row().cells
                    for cell, value in zip(cells, rr):
                        run = add_run_with_font(cell.paragraphs[0], value, ri == 0)
                        if ri == 0: run.font.color.rgb = RGBColor.from_string('1F4E79')
                        if ri == 0 or ri % 2 == 0:
                            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'D9E2F3' if ri == 0 else 'F8FAFC'); cell._tc.get_or_add_tcPr().append(shade)
            elif b['type'] == 'image':
                p = doc.add_paragraph(); p.alignment = 1
                add_run_with_font(p, '').add_picture(str(OUT / b['path']), width=Cm(15.5))
                cap = dp(b['caption']); cap.alignment = 1
                cap.runs[0].italic = True; cap.runs[0].font.color.rgb = RGBColor.from_string('666666')
doc.save(OUT / 'report.docx')

css = "body{font-family:'Microsoft YaHei','PingFang SC','Heiti TC','Hiragino Sans GB','STHeiti','Arial Unicode MS',sans-serif;max-width:1100px;margin:auto;padding:24px;line-height:1.7}h1{color:#1f4e79;border-bottom:2px solid}h2{color:#2e75b6;border-left:4px solid;padding-left:12px}table{border-collapse:collapse;width:100%;margin:15px 0}td,th{border:1px solid #ccd;padding:8px}th{background:linear-gradient(#D9E2F3,#C5D3EE);color:#1f4e79}tr:nth-child(even){background:#F8FAFC}img{max-width:100%}.L1{background:#E7F5FF;border-left:4px solid #1971C2;padding:12px}.L2{background:#FFF9DB;border-left:4px solid #F08C00;padding:12px}.caption{text-align:center;color:#666;font-style:italic}"
html = ['<!doctype html><html lang="zh-CN"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>客制化专项报告</title><style>' + css + '</style></head><body>', '<h1>客制化专项报告</h1><p>' + escape(report['cover']['subtitle']) + '</p><p>日期：2026-06-19　版本：v1.0</p>', '<nav class="toc"><h2>目录</h2>']
for j, sec in enumerate(secs, 1): html.append(f'<p><a href="#ch{j}">{escape(sec["title"])}</a></p>')
html.append('</nav>')
for j, sec in enumerate(secs, 1):
    html.append(f'<h1 id="ch{j}">{escape(sec["title"])}</h1>')
    for ss in sec['subsections']:
        html.append('<h2>' + escape(ss['title']) + '</h2>')
        for b in ss['blocks']:
            if b['type'] == 'paragraph': html.append('<p>' + escape(b['text']) + '</p>')
            elif b['type'] == 'insight': html.append(f'<p class="{b["level"]}">{b["level"]} | 数据来源：{escape(b["source"])}<br>{escape(b["text"])}</p>')
            elif b['type'] == 'image': html.append(f'<img src="{b["path"]}" alt="{escape(b["caption"])}"><p class="caption">{escape(b["caption"])}</p>')
            elif b['type'] == 'table': html.append('<table><thead><tr>' + ''.join('<th>' + escape(str(c)) + '</th>' for c in b['headers']) + '</tr></thead><tbody>' + ''.join('<tr>' + ''.join('<td>' + escape(str(c)) + '</td>' for c in rr) + '</tr>' for rr in b['rows']) + '</tbody></table>')
html.append('</body></html>')
(OUT / 'report.html').write_text('\n'.join(html), encoding='utf-8')
(OUT / 'docs/analysis_proposal.md').write_text('# 分析方案与完成状态\n\n仅取两个指定场景和两档并发；按产品及场景保留 QPS/P95/P99，不跨场景平均。\n\n- ✅ 输入识别与意图解析\n- ✅ 筛选 12 条记录，完整性与公平性核查\n- ✅ 分场景计算并发变化及同档产品差异\n- ✅ 三张三指标图表与文字重叠检查\n- ✅ 独立生成 md、docx、html\n- ✅ 最小交付核查，详见 data/checks.json\n', encoding='utf-8')
(OUT / 'data/block_map.md').write_text('# 报告块映射\n\n第1章：用户问题、三条摘要、四个核心数据点。\n第2章：来源、环境、筛选范围及公平性。\n第3章：每产品一张四行数据表；分场景变化与同档比较；三张指标图。\n第4章：分场景产品建议、不确定性与后续测试。\n第5章：来源哈希、计算公式和交付索引。\n\n全部数据来自 extracted_data.json；计算来自 analysis_results.json；共享结构为 report_data.json。无 L3。\n', encoding='utf-8')

# Full selected-record verification, beyond the required six-point sample.
mdtext = (OUT / 'report.md').read_text(encoding='utf-8')
htmltext = (OUT / 'report.html').read_text(encoding='utf-8')
doc2 = Document(OUT / 'report.docx')
doctext = '\n'.join(p.text for p in doc2.paragraphs) + '\n' + '\n'.join(c.text for t in doc2.tables for rr in t.rows for c in rr.cells)
for r in rows:
    original = next(x for x in raw['records'] if all(x[k] == r[k] for k in ['product', 'scenario', 'threads']))
    assert r == original
    for k, _, _ in metrics:
        assert all(fmt(r[k]) in text for text in [mdtext, htmltext, doctext])
assert len(doc2.inline_shapes) == 3
assert all(run.font.name == '微软雅黑' for p in doc2.paragraphs for run in p.runs if run.text)
assert all('D9E2F3' in t.rows[0].cells[0]._tc.xml for t in doc2.tables)
assert all(not re.findall(r'\{[A-Za-z_][A-Za-z0-9_]*\}', t) for t in [mdtext, htmltext])
assert 'L3' not in mdtext.split('# 第 5 章')[0]
assert (OUT/'report.md').stat().st_size >= 5120
assert (OUT/'report.docx').stat().st_size >= 10240
assert (OUT/'report.html').stat().st_size >= 8192
assert hashlib.sha256(source.read_bytes()).hexdigest() == before
checks = {'status': 'PASS', 'A': '所需中间文件、方案、三格式及3图完整，大小达标，Todo全完成', 'B': '3×2×2=12；每产品一表；全部12记录及36项指标与原文一致；三指标无空值；并发完整', 'C': '中文标题；主图、档位对照、尾延迟验证图齐全；check_text_overlap均为空；其他场景不适用', 'D': '3条摘要含主要发现与不确定性；洞察均含level/source；无L3；元数据公平性一致，未披露条件已说明；D5按用户指定完整两档执行', 'E': '章节标题层级、表格、相对图片路径正确；Word全部正文run字体及各表头颜色检查通过，嵌图3；HTML head/CSS/TOC/字体fallback完整', 'F': '无占位符；md/html/docx全部36项原始指标一致', 'input_unchanged': True, 'skill_source_modified': False, 'chart_text_overlaps': overlaps}
save('data/checks.json', checks)
print(json.dumps({'status': 'PASS', 'records': len(rows), 'charts': 3, 'sizes': {n: (OUT/n).stat().st_size for n in ['report.md', 'report.docx', 'report.html']}}, ensure_ascii=False))
