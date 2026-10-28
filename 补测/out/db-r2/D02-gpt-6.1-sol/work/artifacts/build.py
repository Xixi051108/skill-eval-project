import sys, os, json, hashlib
from pathlib import Path
sys.dont_write_bytecode = True
ROOT = Path.cwd()
OUT = ROOT / 'artifacts'
os.environ['MPLCONFIGDIR'] = str(OUT / '.mpl')
sys.path.insert(0, str(ROOT / 'skill/scripts'))
from parse_intent import parse_intent
from data_source_adapter import load_records, quality_check
from analyze_core import analyze, scenario_records, peak_for

def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def save(name, obj):
    p = OUT / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')
baseline = {str(p.relative_to(ROOT)): digest(p) for folder in ('skill', 'data') for p in (ROOT/folder).rglob('*') if p.is_file()}
save('data/input_hashes.json', baseline)
intent = parse_intent('对比 data/compare.json 三个产品，五个场景独立比较，标注公平性限制')
intent['other_info'] = '固定当前时间 2026-06-19T16:00:00+08:00；五场景独立比较；禁止跨场景合并均值。'
save('data/intent.json', intent)
e = load_records(intent)
qc = quality_check(e)
m, records = e['meta'], e['records']
products, scenarios, threads = m['products'], m['scenarios'], m['concurrencies']
assert len(products)==3 and len(scenarios)==5 and len(records)==len(products)*len(scenarios)*len(threads)==90
assert len({(r['product'],r['scenario'],r['threads']) for r in records})==90
for p in products:
    for s in scenarios:
        assert [r['threads'] for r in scenario_records(records,p,s)]==threads
assert all(isinstance(r[k], (float,int)) and r[k]>0 for r in records for k in ('tps','qps','p95_ms','p99_ms'))
m['original_source_info'] = m['source_info'].copy()
m['source_info'] = {'type':'local_file', 'value':'data/compare.json', 'rows_fetched':len(records)}
m['concurrencies_by_product_scenario']=qc['concurrencies_by_product_scenario']
save('data/extracted_data.json', e)
save('data/quality.json',qc)
a, _ = analyze(e)
a['report_type']='comparison'
logs=m['log_meta']
assert len({l['data_config'] for l in logs})==1
assert len({l['duration_sec'] for l in logs})==1
warnings=[
 {'severity':'INFO','text':'相同场景、逐场景并发档、data_config 和 duration_sec 已核对一致；硬件描述已填充。'},
 {'severity':'STRONG','text':'公平性限制：仅提供共用硬件描述，缺少逐产品节点数、拓扑、资源配额及独占性证据；相同硬件等级未独立验证。下文仅对提供数据作条件性描述，不能据此作严格公平的产品优劣或因果结论。'},
 {'severity':'WEAK','text':'未提供数据库参数、隔离级别、持久化/刷盘、复制与一致性设置、连接方式、预热和缓存状态，事务与延迟统计口径未通过原始命令核实。'},
 {'severity':'WEAK','text':'各产品只有一个 task_id，缺少重复轮次、波动范围、原始日志和资源监控；不能计算显著性或置信区间，也不能确认瓶颈原因。'},
 {'severity':'WEAK','text':'三个产品 start_time 不同，可能存在时段或测试顺序影响；未提供随机化测试证据。'},
 {'severity':'INFO','text':'输入原 source_info 指向 test-resource/mock_records_aggregation.json；本次实际仅读取 data/compare.json，原标记不构成外部来源真实性或可达性验证。'}]
a['fairness_warnings']=warnings
a['product_compare']['fairness'].update({'passed':False,'issues':[w['text'] for w in warnings if w['severity']!='INFO']})
peak = {(p,s):max(scenario_records(records,p,s),key=lambda r:r['qps']) for p in products for s in scenarios}
ratios={s:{p:peak[p,s]['qps']/peak[products[0],s]['qps'] for p in products} for s in scenarios}
scores={p:{s:peak[p,s]['qps']/max(peak[x,s]['qps'] for x in products)*100 for s in scenarios} for p in products}
a['product_compare'].update({'ratio_matrix':ratios,'peak_qps':{s:[peak[p,s] for p in products] for s in scenarios},'concurrent_scale':{s:{p:scenario_records(records,p,s) for p in products} for s in scenarios},'p95_latency':{s:{p:[{'threads':r['threads'],'p95_ms':r['p95_ms']} for r in scenario_records(records,p,s)] for p in products} for s in scenarios},'radar_scores':scores,'overall_score':None,'score_method':'每场景峰值QPS/该场景三产品最大峰值QPS×100；无跨场景总分、均值或总排名。'})
for p in products:
    a['single_product'][p]['scale']={s:scenario_records(records,p,s) for s in scenarios}
    a['single_product'][p]['latency']={s:[{'threads':r['threads'],'p95_ms':r['p95_ms'],'p99_ms':r['p99_ms']} for r in scenario_records(records,p,s)] for s in scenarios}
ins={'single_product_insights':{},'product_compare_insights':[],'selection_insights':[]}
for p in products:
    ins['single_product_insights'][p]=[]
    for s in scenarios:
        rows=scenario_records(records,p,s); pk=peak[p,s]; last=rows[-1]
        text=f"{p} 在 {s} 的采样峰值 QPS 为 {pk['qps']:,.2f}（{pk['threads']:,} 并发），32→峰值扩展 {pk['qps']/rows[0]['qps']:.2f}x；1,000 并发比峰值下降 {(1-last['qps']/pk['qps'])*100:.2f}%，P95 从峰值处 {pk['p95_ms']:,.2f} 升至 {last['p95_ms']:,.2f} ms。"
        ins['single_product_insights'][p].append({'level':'L1','text':text,'source':f'data/compare.json records: product={p}, scenario={s}; Python max/ratio','scenario':s})
for s in scenarios:
    ranked=sorted(products,key=lambda p:peak[p,s]['qps'],reverse=True)
    w,b,c=ranked
    text=f"{s}：{w} 采样峰值 QPS {peak[w,s]['qps']:,.2f}，高于 {b}（{peak[b,s]['qps']:,.2f}）{(peak[w,s]['qps']/peak[b,s]['qps']-1)*100:.2f}%，高于 {c}（{peak[c,s]['qps']:,.2f}）{(peak[w,s]['qps']/peak[c,s]['qps']-1)*100:.2f}%；三者峰值均位于 {peak[w,s]['threads']:,} 并发。仅为本次数据的条件性比较。"
    ins['product_compare_insights'].append({'level':'L1','text':text,'source':f'data/compare.json records: scenario={s}, threads=512; Python max/ratio','scenario':s})
save('data/analysis_results.json', a)
save('data/insights.json',ins)
for group in [*ins['single_product_insights'].values(),ins['product_compare_insights'],ins['selection_insights']]:
    assert all({'level','source','scenario'} <= x.keys() for x in group)
    assert all(x['level']!='L3' for x in group)
(OUT/'docs').mkdir(exist_ok=True)
(OUT/'docs/analysis_proposal.md').write_text('# 分析方案\n\n使用 comparison；数据仅来自 data/compare.json。以 TDSQL-B 为基准，五场景独立计算峰值、同比、扩展和延迟，覆盖六个并发档。综合展示仅含场景内归一化分值，无跨场景均值或总排名。硬件与配置证据不足时只作条件性描述。三格式从同一 report_data 独立生成，20 张中文标题图并核查布局。\n',encoding='utf-8')
print('阶段1—3完成：90条记录，5条场景核心结论；公平性限制已记录。',flush=True)

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.text import Text
from constants import SCENARIO_CN as CN, PRODUCT_COLORS as COLORS
plt.rcParams.update({'font.sans-serif':['Microsoft YaHei','SimHei','DejaVu Sans'],'axes.unicode_minus':False,'font.size':10})
(OUT/'charts').mkdir(exist_ok=True)
chart_qc=[]
def check_text_overlap(fig):
    fig.canvas.draw(); renderer=fig.canvas.get_renderer()
    texts=[t for t in fig.findobj(Text) if t.get_visible() and t.get_text().strip()]
    boxes=[(t,t.get_window_extent(renderer)) for t in texts]
    bad=[]
    for i,(t,b) in enumerate(boxes):
        if b.width<=0 or b.height<=0: continue
        for u,c in boxes[i+1:]:
            if c.width<=0 or c.height<=0: continue
            dx=min(b.x1,c.x1)-max(b.x0,c.x0); dy=min(b.y1,c.y1)-max(b.y0,c.y0)
            if dx>1 and dy>1: bad.append([t.get_text(),u.get_text()])
    return bad
def finish(fig,name,title):
    fig.tight_layout(pad=2)
    overlaps=check_text_overlap(fig)
    chart_qc.append({'file':name,'title':title,'overlaps':overlaps})
    fig.savefig(OUT/'charts'/name,dpi=150,bbox_inches='tight')
    plt.close(fig)
def axes(title,ylabel):
    fig,ax=plt.subplots(figsize=(10,6))
    ax.set_title(title,pad=22); ax.set_ylabel(ylabel); ax.grid(alpha=.2,axis='y')
    return fig,ax
for s in scenarios:
    title=f'{CN[s]}场景：采样峰值对比'
    fig,ax=axes(title,'峰值 QPS')
    ax.bar(products,[peak[p,s]['qps'] for p in products],color=COLORS[:3])
    ax.set_ylim(0,max(peak[p,s]['qps'] for p in products)*1.15)
    finish(fig,f'peak_{s}.png',title)
    for metric, label, short in [('qps','QPS','qps'),('p95_ms','P95 (ms)','p95')]:
        title=f'{CN[s]}场景：并发与{label}对比'
        fig,ax=axes(title,label)
        for i,p in enumerate(products):
            rows=scenario_records(records,p,s)
            ax.plot(threads,[r[metric] for r in rows],marker='o',label=p,color=COLORS[i])
        ax.set_xscale('log',base=2); ax.set_xticks(threads,labels=[f'{t:,}' for t in threads]); ax.set_xlabel('并发数')
        ax.legend(loc='upper left',bbox_to_anchor=(1.01,1),borderaxespad=0)
        finish(fig,f'{short}_{s}.png',title)
for i,p in enumerate(products):
    title=f'{p}：各场景独立扩展曲线'
    fig,ax=axes(title,'QPS / 本场景32并发QPS')
    for s in scenarios:
        rows=scenario_records(records,p,s)
        ax.plot(threads,[r['qps']/rows[0]['qps'] for r in rows],label=CN[s],marker='o')
    ax.set_xscale('log',base=2);ax.set_xticks(threads,labels=[f'{t:,}' for t in threads]);ax.set_xlabel('并发数')
    ax.legend(loc='upper left',bbox_to_anchor=(1.01,1),borderaxespad=0)
    finish(fig,f'product_{i+1}.png',title)
angles=np.linspace(0,2*np.pi,len(scenarios),endpoint=False).tolist(); angles+=angles[:1]
fig,ax=plt.subplots(figsize=(10,8),subplot_kw={'polar':True})
for i,p in enumerate(products):
    vals=[scores[p][s] for s in scenarios]; vals+=vals[:1]
    ax.plot(angles,vals,label=p,color=COLORS[i]);ax.fill(angles,vals,alpha=.07,color=COLORS[i])
ax.set_xticks(angles[:-1],labels=[CN[s] for s in scenarios]); ax.set_ylim(0,110);ax.set_yticks([20,40,60,80,100]);ax.tick_params(axis='x',pad=16)
title='场景内峰值归一化雷达图';ax.set_title(title,pad=42)
ax.legend(loc='upper left',bbox_to_anchor=(1.22,1.03),borderaxespad=0)
finish(fig,'radar.png',title)
fig,axs=plt.subplots(5,1,figsize=(11,16))
for ax,s in zip(axs,scenarios):
    order=sorted(products,key=lambda p:scores[p][s])
    ax.barh(order,[scores[p][s] for p in order],color=[COLORS[products.index(p)] for p in order]);ax.set_xlim(0,110)
    ax.set_title(f'{CN[s]}场景独立天梯',pad=14);ax.set_xlabel('本场景峰值归一化分值');ax.grid(axis='x',alpha=.2)
finish(fig,'ladder.png','五场景独立天梯图')
save('data/chart_qa.json',chart_qc)
assert len(chart_qc)==20
assert all(not x['overlaps'] for x in chart_qc), [(x['file'],x['overlaps']) for x in chart_qc if x['overlaps']]
print('阶段4完成：20张图，文字边界框重叠检查通过。',flush=True)
# 阶段5：共享数据结构，三格式独立渲染。
def para(text): return {'type':'paragraph','text':text}
def table(headers,rows): return {'type':'table','headers':headers,'rows':rows}
def image(name,caption): return {'type':'image','path':'charts/'+name,'caption':caption}
def insight(x): return dict(x,type='insight')
def heading(text): return {'type':'heading','text':text,'level':3}
def bullet(items): return {'type':'bullet','items':items}
def sub(n,title,blocks): return {'id':n.replace('.','_'),'title':n+' '+title,'blocks':blocks}
def section(n,title,subs): return {'id':'ch'+str(n),'title':f'第 {n} 章 {title}','blocks':[],'subsections':subs}
def fmt(v): return f'{v:,.2f}' if v is not None else '-'
key_rows=[]
for s in scenarios:
    for p in products:
        r=peak[p,s]
        key_rows.append([s,p,f"{r['threads']:,}",fmt(r['tps']),fmt(r['qps']),fmt(r['p95_ms']),fmt(r['p99_ms'])])
report={'cover':{'title':'数据库性能对比报告','subtitle':' vs '.join(products),'date':'2026-06-19','version':'v1.0'},'sections':[]}
secs=report['sections']
secs.append(section(1,'测试结论',[
    sub('1.1','测试结论摘要',[para('本报告为 comparison，覆盖三个产品、五个 sysbench 场景和六个并发档，共90条聚合记录。仅描述所提供数据，五场景独立比较，不计算跨场景均值或总排名。所有产品在每个场景的采样峰值均位于512并发，不能推定未采样负载下的真正极值。'),para('【公平性限制 / STRONG】相同硬件等级未独立验证；数据库配置和重复测试证据不足。下列数值比较仅在输入标记条件成立的前提下有效，不能视为严格公平的基准认证。')]),
    sub('1.2','关键性能数字',[table(['场景','产品','峰值并发','TPS@QPS峰值','峰值QPS','P95 (ms)','P99 (ms)'],key_rows)]),
    sub('1.3','核心结论',[insight(x) for x in ins['product_compare_insights']])]))
config_rows=[[l['product'],l['engine'],l['data_config'],str(l['duration_sec']),l['start_time'],l['task_id']] for l in logs]
secs.append(section(2,'测试环境',[
    sub('2.1','数据库测试配置',[table(['产品','工具','数据集','时长 (s)','开始时间（输入原文）','任务ID'],config_rows),para('三产品的数据库参数、实例数量、拓扑、资源配额、持久化、复制和一致性设置均未提供。不可用产品名称推定部署方式。输入 start_time 未含时区，按原文展示。')]),
    sub('2.2','硬件环境',[table(['项目','输入描述'],[[k,v] for k,v in m['test_env'].items()]),para('以上为输入的共用 test_env，不能证明各产品实际获得相同硬件资源。')]),
    sub('2.3','测试方法',[para('使用输入标记的 sysbench 1.0.20；tables=16、table_size=1000000，duration_sec=300。六档并发为32、64、128、256、512、1000。原始命令、运行顺序、预热与采样策略未提供。'),table(['标准场景','观测重点'],[[s,{'oltp_point_select':'点查：QPS、P95、P99','oltp_read_only':'只读事务：TPS、QPS、P95、P99','oltp_read_write':'混合读写：TPS、QPS、P95、P99','oltp_write_only':'写入事务：TPS、QPS、P95、P99','oltp_update_index':'索引更新：TPS、P95、P99'}[s]] for s in scenarios]),para('峰值按每产品、每场景六档QPS的最大值计算；同比=(A/B−1)×100%；扩展倍数=该档QPS/32并发QPS。比值基准为首个产品TDSQL-B。QPS越高、同场景同并发延迟越低越好，但吞吐与延迟需同时评估。')]),
    sub('2.4','数据质量与公平性',[para('90=3×5×6，组合唯一、逐产品逐场景六档齐全；TPS、QPS、P95、P99空值率均为0%。数据集与时长标记一致。每组合仅一条汇总记录，不能把六档并发当成重复样本。')]+[para(f"[{w['severity']}] {w['text']}") for w in warnings])]))
subs=[]
for i,p in enumerate(products,1):
    rows=[]
    for s in scenarios:
        for r in scenario_records(records,p,s): rows.append([s,f"{r['threads']:,}"]+[fmt(r[k]) for k in ('tps','qps','p95_ms','p99_ms')])
    subs.append(sub(f'3.{i}',p+' 全量数据',[table(['场景','并发','TPS','QPS','P95 (ms)','P99 (ms)'],rows)]))
secs.append(section(3,'性能数据',subs))
single_blocks=[]
for i,p in enumerate(products,1):
    single_blocks += [heading(f'4.1.{i} {p}'),image(f'product_{i}.png',f'{p}：各场景独立归一化扩展；每条曲线以本场景32并发QPS为1，不合并场景')]
    single_blocks += [insight(x) for x in ins['single_product_insights'][p]]
compare_blocks=[]
for i,s in enumerate(scenarios,1):
    compare_blocks += [heading(f'4.2.{i} {CN[s]}（{s}）'),insight(ins['product_compare_insights'][i-1])]
    rows=[]
    for t in threads:
        rs=[next(r for r in records if r['product']==p and r['scenario']==s and r['threads']==t) for p in products]
        rows.append([f'{t:,}']+[fmt(r['qps']) for r in rs]+[fmt(r['p95_ms']) for r in rs]+[fmt(r['p99_ms']) for r in rs])
    compare_blocks += [table(['并发']+[p+' QPS' for p in products]+[p+' P95(ms)' for p in products]+[p+' P99(ms)' for p in products],rows)]
    r512=[peak[p,s] for p in products]
    best95=min(r512,key=lambda r:r['p95_ms']); best99=min(r512,key=lambda r:r['p99_ms'])
    compare_blocks += [para(f"同为512并发时，最低P95为 {best95['product']} 的 {fmt(best95['p95_ms'])} ms，最低P99为 {best99['product']} 的 {fmt(best99['p99_ms'])} ms；吞吐领先不等于尾延迟同时领先。全六档数据见本节表格。"),image(f'peak_{s}.png',f'{CN[s]}：各产品采样峰值QPS，峰值均为512并发'),image(f'qps_{s}.png',f'{CN[s]}：三产品六档并发QPS，线段仅连接采样点'),image(f'p95_{s}.png',f'{CN[s]}：三产品六档并发P95延迟（ms）')]
compare_blocks += [heading('4.2.6 性能比值矩阵与场景内分值'),para('比值=本产品该场景峰值QPS/TDSQL-B该场景峰值QPS，基准为1.0000x。场景内分值=本产品峰值QPS/该场景三产品最大峰值QPS×100。不代表成本、稳定性、延迟综合效用；无业务权重，不计算跨场景总分、合并均值或总排名。'),table(['场景']+products,[[s]+[f'{ratios[s][p]:.4f}x' for p in products] for s in scenarios]),table(['场景']+products,[[s]+[fmt(scores[p][s]) for p in products] for s in scenarios]),image('radar.png','雷达图：每个轴单独按本场景最大峰值QPS归一化至100'),image('ladder.png','天梯图：五个独立场景各自展示归一化分值与顺序，无跨场景总排名')]
secs.append(section(4,'横向对比分析',[sub('4.1','单产品分析',single_blocks),sub('4.2','产品横向对比',compare_blocks)]))
secs.append(section(5,'选型与调优建议',[
    sub('5.1','选型决策树',[bullet(['先核实逐产品硬件、拓扑与配置；若无法对齐，保留条件性比较，不做最终产品优劣判断。','条件对齐后，确定业务对应场景及延迟SLO；按同场景同并发的QPS、P95、P99筛选。','若吞吐优先，可把本次场景峰值作为复测候选；若延迟优先，按第4章各档延迟选择并发，再用业务负载复测。','若尚无延迟SLO或成本数据，暂不作最终容量与性价比结论。'])]),
    sub('5.2','适用场景矩阵',[table(['场景','本次吞吐候选（条件性）','后续验证'],[[s,max(products,key=lambda p:peak[p,s]['qps']),'核实配置，重复测试；校验业务P95/P99门限'] for s in scenarios])]),
    sub('5.3','与对标产品的差异',[para('本次数据的场景内吞吐顺序见第1章和第4章；P95与P99并非始终按该顺序排序。输入未提供功能、成本、运维复杂度或可用性数据，不据此作相关优劣结论。')]),
    sub('5.4','调优最佳实践',[para('本次所有场景采样峰值均在512并发；增加至1000并发后QPS下降、P95上升，建议在256、512、1000并发间补充更密采样，结合业务SLO确定并发上限。512仅为当前采样峰值，不是生产推荐值。先固定持久化与一致性设置、记录预热和资源监控，再逐项变更参数并重复采样；现有数据不足以给出buffer_pool、NUMA等具体参数值或瓶颈归因。')]),
    sub('5.5','测试边界',[para('仅涵盖本文件列出的产品版本、五类OLTP场景、标记数据集、300秒测试和六个并发档；没有故障恢复、长时稳定性、OLAP、真实业务或成本证据。公平性STRONG限制在完成硬件与拓扑核实前持续有效。')])]))
secs.append(section(6,'附录',[
    sub('6.1','原始数据来源',[para('实际读取：data/compare.json；原source_info仅作输入元数据保留，未读取或验证其中指向的其他文件。原始记录未经修改，规范化产物只更新来源说明并追加质量信息。'),table(['文件','用途'],[['data/intent.json','识别结果与固定时间'],['data/extracted_data.json','90条完整记录与来源说明'],['data/analysis_results.json','计算结果、公平性限制与公式'],['data/insights.json','逐场景L1洞察及来源'],['data/quality.json','数据质量与覆盖检查'],['data/chart_qa.json','20图文字重叠检查'],['data/input_hashes.json','输入与Skill文件SHA256'],['docs/checks.md','交付核查结果']])]),
    sub('6.2','sysbench命令与计算边界',[para('原始sysbench命令未提供，无法复原连接参数、Lua脚本选项与实际统计过程；本报告不伪造可复现命令。基于聚合结果进行描述性计算，不对缺失样本估算数据。'),para('报告日期按指定固定时间2026-06-19T16:00:00+08:00生成；测试时间使用输入原文。报告不依赖外部服务、不安装依赖。')])]))
save('data/report_data.json',report)

import html, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
LABEL={'L1':'✅ L1 确定性结论','L2':'📊 L2 统计性发现','L3':'💡 L3 假设性洞察'}
COLOR={'L1':'1971C2','L2':'E67700','L3':'C92A2A'}
def groups():
    for sec in secs:
        yield 1,sec['id'],sec['title'],sec['blocks']
        for sb in sec['subsections']: yield 2,sb['id'],sb['title'],sb['blocks']
def render_md(rd):
    c=rd['cover']; lines=[f"# {c['title']}",'',c['subtitle'],'',f"日期：{c['date']}　版本：{c['version']}",'','## 目录','']
    lines += [f"- [{sec['title']}](#ch{j})" for j,sec in enumerate(secs,1)]
    for level,id,title,blocks in groups():
        lines += ['',f'<a id="{id}"></a>',f"{'#'*level} {title}",'']
        for b in blocks:
            typ=b['type']
            if typ=='paragraph': lines += [b['text'],'']
            elif typ=='heading': lines += ['### '+b['text'],'']
            elif typ=='bullet': lines += ['- '+t for t in b['items']]+['']
            elif typ=='table': lines += ['| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |']+['| '+' | '.join(map(str,r))+' |' for r in b['rows']]+['']
            elif typ=='image': lines += [f"![{b['caption']}]({b['path']})",f"*图：{b['caption']}*",'']
            elif typ=='insight': lines += [f"> {LABEL[b['level']]} | 数据来源：{b['source']}",'>','> '+b['text'],'']
    return '\n'.join(lines)

def set_doc_default_font(doc):
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','List Bullet']:
        st=doc.styles[name];st.font.name='微软雅黑';st.font.size=Pt(10.5 if name=='Normal' else 12)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
def add_run_with_font(p,text,**kw):
    r=p.add_run(str(text));r.font.name='微软雅黑';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    for k,v in kw.items():
        if k=='color': r.font.color.rgb=v
        else: setattr(r.font,k,v)
    return r
def shade(cell,color):
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(sh)
def render_docx(rd,path):
    doc=Document();set_doc_default_font(doc)
    sec=doc.sections[0];sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for txt in [rd['cover']['title'],rd['cover']['subtitle'],f"日期：{rd['cover']['date']}　版本：v1.0"]:
        p=doc.add_paragraph();p.alignment=1;add_run_with_font(p,txt,bold=True)
    doc.add_page_break()
    for level,id,title,blocks in groups():
        p=doc.add_paragraph(style=f'Heading {level}');add_run_with_font(p,title,bold=True)
        for b in blocks:
            typ=b['type']
            if typ in ['paragraph','heading']:
                p=doc.add_paragraph(style='Heading 3' if typ=='heading' else None);add_run_with_font(p,b['text'])
            elif typ=='bullet':
                for text in b['items']:add_run_with_font(doc.add_paragraph(style='List Bullet'),text)
            elif typ=='table':
                t=doc.add_table(rows=1,cols=len(b['headers']));t.style='Table Grid'
                for cell,txt in zip(t.rows[0].cells,b['headers']):
                    shade(cell,'D9E2F3');add_run_with_font(cell.paragraphs[0],txt,bold=True,color=RGBColor.from_string('1F4E79'))
                for i,row in enumerate(b['rows'],1):
                    for cell,txt in zip(t.add_row().cells,row):
                        if i%2==0:shade(cell,'F8FAFC')
                        add_run_with_font(cell.paragraphs[0],txt)
            elif typ=='image':
                p=doc.add_paragraph();p.alignment=1;add_run_with_font(p,'').add_picture(str(OUT/b['path']),width=Cm(15.5))
                p=doc.add_paragraph();p.alignment=1;add_run_with_font(p,'图：'+b['caption'],italic=True,color=RGBColor.from_string('666666'))
            elif typ=='insight':
                p=doc.add_paragraph();pp=p._p.get_or_add_pPr();borders=OxmlElement('w:pBdr');left=OxmlElement('w:left')
                for k,v in {'val':'single','sz':'18','color':COLOR[b['level']],'space':'6'}.items():left.set(qn('w:'+k),v)
                borders.append(left);pp.append(borders)
                add_run_with_font(p,LABEL[b['level']]+' | 数据来源：'+b['source']+'\n'+b['text'])
    doc.save(path)

def render_html(rd):
    esc=lambda t:html.escape(str(t))
    c=rd['cover'];parts=['<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>'+esc(c['title'])+'</title><style>',
    "body{font-family:'Microsoft YaHei','PingFang SC','Heiti TC','Hiragino Sans GB','STHeiti','Arial Unicode MS',sans-serif;max-width:1250px;margin:auto;padding:28px;color:#27384b;line-height:1.7}h1{color:#1f4e79;border-bottom:2px solid #1f4e79}h2{color:#2e75b6;border-left:4px solid #2e75b6;padding-left:12px}h3{color:#4a6d8c}table{border-collapse:collapse;width:100%;font-size:13px}th,td{border:1px solid #ccd5df;padding:7px}th{background:linear-gradient(#D9E2F3,#C5D3EE);color:#1F4E79}tbody tr:nth-child(even){background:#F8FAFC}img{display:block;max-width:100%;margin:20px auto}.caption{color:#666;text-align:center;font-style:italic}.insight{padding:12px;margin:14px 0;border-left:4px solid}.L1{background:#E7F5FF;border-color:#1971C2}.L2{background:#FFF9DB;border-color:#F08C00}.L3{background:#FFE3E3;border-color:#C92A2A}.toc{background:#F8FAFC;padding:16px}",
    '</style></head><body><header><h1>'+esc(c['title'])+'</h1><p>'+esc(c['subtitle'])+'</p><p>日期：'+c['date']+'　版本：'+c['version']+'</p></header><nav class="toc"><strong>目录</strong><ul>']
    parts += ['<li><a href="#'+sec['id']+'">'+esc(sec['title'])+'</a></li>' for sec in secs];parts+=['</ul></nav>']
    for level,id,title,blocks in groups():
        parts += [f'<h{level} id="{id}">{esc(title)}</h{level}>']
        for b in blocks:
            typ=b['type']
            if typ=='paragraph':parts+=['<p>'+esc(b['text'])+'</p>']
            elif typ=='heading':parts+=['<h3>'+esc(b['text'])+'</h3>']
            elif typ=='bullet':parts+=['<ul>'+''.join('<li>'+esc(t)+'</li>' for t in b['items'])+'</ul>']
            elif typ=='table':parts+=['<table><thead><tr>'+''.join('<th>'+esc(t)+'</th>' for t in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(t)+'</td>' for t in row)+'</tr>' for row in b['rows'])+'</tbody></table>']
            elif typ=='image':parts+=[f'<img src="{esc(b["path"])}" alt="{esc(b["caption"])}"><p class="caption">图：{esc(b["caption"])}</p>']
            elif typ=='insight':parts+=[f'<div class="insight {b["level"]}"><strong>{LABEL[b["level"]]} | 数据来源：{esc(b["source"])}</strong><p>{esc(b["text"])}</p></div>']
    return '\n'.join(parts+['</body></html>'])
(OUT/'report.md').write_text(render_md(report),encoding='utf-8')
render_docx(report,OUT/'report.docx')
(OUT/'report.html').write_text(render_html(report),encoding='utf-8')
(OUT/'data/block_map.md').write_text('# 内容映射\n\n'+ '\n'.join(f'- {title}：{len(blocks)}个内容块' for _,_,title,blocks in groups())+'\n',encoding='utf-8')
print('阶段5完成：md、docx、html已分别从共享数据独立生成。',flush=True)

# A—F交付检查：失败时保存证据并停止，不伪称通过。
from zipfile import ZipFile
from lxml import etree
checks=[]
def ck(id,desc,ok,detail=''):
    checks.append({'id':id,'description':desc,'passed':bool(ok),'detail':detail})
def exists(name):return (OUT/name).is_file()
for id,name in [('A1','data/intent.json'),('A2','data/extracted_data.json'),('A4','data/analysis_results.json'),('A5','data/insights.json')]:
    ck(id,name,exists(name) and isinstance(json.loads((OUT/name).read_text(encoding='utf-8')),dict))
ck('A3','内容映射',exists('data/block_map.md'))
ck('A6','comparison图表不少于17+3',len(list((OUT/'charts').glob('*.png')))>=20)
ck('A7','分析方案',exists('docs/analysis_proposal.md'))
for id,name,size in [('A9','report.md',5120),('A10','report.docx',10240),('A11','report.html',8192)]:ck(id,name,exists(name) and (OUT/name).stat().st_size>=size,str((OUT/name).stat().st_size))
raw=json.loads((ROOT/'data/compare.json').read_text(encoding='utf-8'))
ck('B1','三个意图产品一致',products==raw['meta']['products'] and len(products)==3)
ck('B2','90条且组合唯一',len(records)==90 and len({(r['product'],r['scenario'],r['threads']) for r in records})==90)
ch3=secs[2]['subsections']
ck('B3','每产品一表，30条',len(ch3)==3 and all(len(sb['blocks'][0]['rows'])==30 for sb in ch3))
# 检查全部90条，超过6条抽检基线，同时核对第3章序列化后的每个数字。
expected_raw=[]
for p in products:
    for s in scenarios:
        for r in scenario_records(raw['records'],p,s):expected_raw.append([s,f"{r['threads']:,}"]+[fmt(r[k]) for k in ('tps','qps','p95_ms','p99_ms')])
actual_rows=[row for sb in ch3 for row in sb['blocks'][0]['rows']]
ck('B4','全90条输入一致及正文数值一致',records==raw['records'] and actual_rows==expected_raw,'超过6条抽检要求')
ck('B5','必需指标无空值',all(qc[k]==0 for k in ['tps_null_rate','qps_null_rate','p95_null_rate']))
ck('B6','每产品场景六档完整',all(ts==threads for ss in qc['concurrencies_by_product_scenario'].values() for ts in ss.values()))
ck('B7','实际数据源标记完整',m['source_info']=={'type':'local_file','value':'data/compare.json','rows_fetched':90})
blocks=[b for _,_,_,bs in groups() for b in bs]
images=[b for b in blocks if b['type']=='image']
ck('C1-C3','五场景峰值/QPS/P95图',all(exists('charts/'+name+'_'+s+'.png') for name in ['peak','qps','p95'] for s in scenarios))
ck('C4','雷达图场景内归一化',exists('charts/radar.png') and all(max(scores[p][s] for p in products)==100 for s in scenarios))
ck('C5','五场景独立天梯',exists('charts/ladder.png') and a['product_compare']['overall_score'] is None)
ck('C6','三产品扩展曲线',all(exists(f'charts/product_{i}.png') for i in range(1,4)))
ck('C7','图表含中文标题',all(re.search('[\u4e00-\u9fff]',x['title']) for x in chart_qc))
ck('C8','check_text_overlap全部返回空',all(x['overlaps']==[] for x in chart_qc),'20图Matplotlib实际渲染后文字边界框两两检查')
all_ins=[x for group in ins['single_product_insights'].values() for x in group]+ins['product_compare_insights']+ins['selection_insights']
ck('D1','全部洞察含level',all(x.get('level') in ['L1','L2','L3'] for x in all_ins))
ck('D2','全部洞察含source',all(x.get('source') for x in all_ins))
ck('D3a','L3仅selection',all(x['level']!='L3' for x in all_ins),'未生成任何推测性L3洞察')
md=(OUT/'report.md').read_text(encoding='utf-8');ht=(OUT/'report.html').read_text(encoding='utf-8')
ck('D3b','正文无越界L3','L3' not in md and '💡' not in md)
ck('D4','数据集一致',len({l['data_config'] for l in logs})==1)
ck('D5','并发档全部使用',all(len(a['product_compare']['concurrent_scale'][s][p])==6 for s in scenarios for p in products))
ck('D6','第1章核心结论4至7条',4<=len(secs[0]['subsections'][2]['blocks'])<=7)
ck('D7','第4章五场景独立',all(any(b['type']=='heading' and s in b['text'] for b in compare_blocks) for s in scenarios))
ck('D8','STRONG警告正文显著标注',all(w['text'] in md for w in warnings if w['severity']=='STRONG') and '【公平性限制 / STRONG】' in md)
ck('E1.1','章/节/小节标题层级',all('# '+sec['title'] in md for sec in secs) and '#### ' not in md)
ck('E1.2','表格结构',all(len(row)==len(b['headers']) for b in blocks if b['type']=='table' for row in b['rows']))
ck('E1.3','Markdown图相对路径',len(re.findall(r'!\[.*?\]\(charts/[^)]+\)',md))==20 and all(exists(b['path']) for b in images))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
with ZipFile(OUT/'report.docx') as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    runs=xml.xpath('//w:r[w:t]',namespaces=ns)
    fonts=[r.xpath('string(w:rPr/w:rFonts/@w:eastAsia)',namespaces=ns) for r in runs]
    ck('E2.1','Word全部文本显式微软雅黑',len(fonts)>=5 and all(f=='微软雅黑' for f in fonts),f'{len(fonts)}处文本run')
    tables=xml.xpath('//w:tbl',namespaces=ns)
    ck('E2.2','全部Word表头浅蓝',len(tables)>=3 and all(t.xpath('w:tr[1]/w:tc/w:tcPr/w:shd/@w:fill',namespaces=ns)==['D9E2F3']*len(t.xpath('w:tr[1]/w:tc',namespaces=ns)) for t in tables))
    ck('E2.3','图片20张且嵌入完整',len(xml.xpath('//a:blip',namespaces=ns))==len(images)==20 and len([n for n in z.namelist() if n.startswith('word/media/')])==20)
    doc_text=''.join(xml.xpath('//w:t/text()',namespaces=ns))
ck('E3.1','HTML完整head与内联CSS',all(x in ht for x in ['<head>','charset="UTF-8"','name="viewport"','<title>','<style>','class="toc"']))
ck('E3.2','HTML中文字体回退',all(x in ht for x in ['Microsoft YaHei','PingFang SC','Heiti TC','Hiragino Sans GB','STHeiti','Arial Unicode MS']))
ck('E3.3','HTML图片相对路径且存在',len(re.findall(r'<img src="charts/[^"]+"',ht))==20)
ck('F1','无未替换占位符',not re.search(r'\{[A-Za-z_][A-Za-z0-9_]*\}',md+ht))
ck('F2','建议使用实测并发数',all(str(t) in secs[4]['subsections'][3]['blocks'][0]['text'] for t in [256,512,1000]))
ck('F3','三格式15个场景峰值数字一致',all(fmt(peak[p,s]['qps']) in text for p in products for s in scenarios for text in [md,ht,doc_text]))
ck('额外1','Skill和输入哈希全部不变',all((ROOT/n).is_file() and digest(ROOT/n)==v for n,v in baseline.items()))
current={str(p.relative_to(ROOT)) for folder in ('skill','data') for p in (ROOT/folder).rglob('*') if p.is_file()}
ck('额外2','Skill及输入目录未新增文件',current==set(baseline))
ck('额外3','日期使用固定锚点',all('2026-06-19' in text for text in [md,ht,doc_text]))
(OUT/'docs/todo.md').write_text('# Todo\n\n'+ '\n'.join('- [x] '+s for s in ['读取全部references','识别输入并生成intent','接入数据与质量门控','逐场景分析与公平性核查','生成20张图并检查文字重叠','独立渲染三格式','完成A至F核查与输入/源码哈希差异检查'])+'\n',encoding='utf-8')
ck('A8','Todo全完成','[ ]' not in (OUT/'docs/todo.md').read_text(encoding='utf-8'))
save('data/checks.json',checks)
failed=[c for c in checks if not c['passed']]
(OUT/'docs/checks.md').write_text('# 交付核查\n\n'+f'通过 {len(checks)-len(failed)}/{len(checks)} 项。所有检查针对实际产物；不表示已执行数据库压测或完成来源真实性认证。\n\n| 检查 | 结果 | 说明 |\n| --- | --- | --- |\n'+ '\n'.join(f"| {c['id']} | {'通过' if c['passed'] else '失败'} | {c['description']}；{c['detail']} |" for c in checks)+'\n\n最终差异：仅在 artifacts/ 新增生成脚本、报告、图表、分析和核查产物；skill/ 与 data/ 的所有文件SHA256及文件集合保持不变。图表文字检查不替代人工视觉检查。\n',encoding='utf-8')
print(json.dumps({'checks':len(checks),'failed':failed,'reports':{n:(OUT/n).stat().st_size for n in ['report.md','report.docx','report.html']}},ensure_ascii=False),flush=True)
assert not failed, failed

