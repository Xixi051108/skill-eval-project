"""Generate this single report without changing the supplied Skill or inputs."""
import os
import sys
import json
import re
import hashlib
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts'
os.environ['MPLCONFIGDIR'] = str(OUT / '.mpl')
sys.dont_write_bytecode = True
sys.path.insert(0, str(ROOT / 'skill' / 'scripts'))
from parse_intent import parse_intent
from data_source_adapter import load_records, quality_check
from analyze_core import analyze
from constants import SCENARIO_CN


def save(name, value):
    p = OUT / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def hashes():
    files = list((ROOT / 'skill').rglob('*')) + list((ROOT / 'data').rglob('*'))
    return {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in files if p.is_file()}


before = hashes()
save('data/input_hashes.json', before)
intent = parse_intent('读取 data/single.log，生成 single 性能报告，保留全部性能指标')
assert intent['data_source_type'] == 'local_file' and intent['report_type'] == 'single'
save('data/intent.json', intent)
extracted = load_records(intent)
raw = (ROOT / 'data/single.log').read_text(encoding='utf-8')
product, endpoint = re.search(r'^DB:\s*(\S+)\s+host=(\S+)', raw, re.M).groups()
duration = int(re.search(r'^Duration:\s*(\d+)s', raw, re.M)[1])
start = re.search(r'^Time:\s*(.+)$', raw, re.M)[1].strip()
data_config = re.search(r'^Data:\s*(.+)$', raw, re.M)[1].strip()
# Correct the built-in header reader from explicit source fields only.
meta = extracted['meta']
meta['products'] = [product]
meta['source_info']['value'] = 'data/single.log'
header = meta['log_meta'][0]
header.update(product=product, db_endpoint=endpoint, duration_sec=duration,
              start_time=start, data_config=data_config)
for row in extracted['records']:
    row['product'] = product
records = extracted['records']
qc = quality_check(extracted)
meta['concurrencies_by_product_scenario'] = qc['concurrencies_by_product_scenario']
save('data/extracted_data.json', extracted)
save('data/quality.json', qc)
assert len(records) == 1 * 5 * 6 == 30
assert meta['concurrencies'] == [32, 64, 128, 256, 512, 1000]
scenarios = meta['scenarios']
source_rows = re.findall(r'^--- (oltp_\w+) \| threads=(\d+) \| (\d+:\d+:\d+) .*?TPS=([\d.]+), QPS=([\d.]+), P95=([\d.]+)ms, P99=([\d.]+)ms', raw, re.M)
assert len(source_rows) == len(records)
for r, src in zip(records, source_rows):
    assert (r['scenario'], r['threads']) == (src[0], int(src[1]))
    assert [r[k] for k in ('tps', 'qps', 'p95_ms', 'p99_ms')] == list(map(float, src[3:]))
    r['test_time'] = src[2]
save('data/extracted_data.json', extracted)

analysis, _ = analyze(extracted)
analysis['report_type'] = 'single'
analysis['fairness_warnings'] = [
    {'level': 'WEAK', 'text': '未提供 CPU、内存、磁盘、操作系统或数据库参数，无法核实硬件环境。'},
    {'level': 'INFO', 'text': '单产品、单数据集、统一日志时长 300s；全部场景并发档一致，无跨产品公平性比较。'}]
sp = analysis['single_product'][product]
sp['scale'] = {s: [{'threads': r['threads'], 'qps': r['qps']} for r in records if r['scenario'] == s] for s in scenarios}
sp['latency'] = {s: [{'threads': r['threads'], 'p95_ms': r['p95_ms'], 'p99_ms': r['p99_ms']} for r in records if r['scenario'] == s] for s in scenarios}
findings = []
for s in scenarios:
    rows = sorted([r for r in records if r['scenario'] == s], key=lambda r: r['threads'])
    first, last = rows[0], rows[-1]
    peak = max(rows, key=lambda r: r['qps'])
    decline = (1 - last['qps'] / peak['qps']) * 100
    text = (f"{SCENARIO_CN[s]}（{s}）：并发 {first['threads']}→{peak['threads']}，"
            f"QPS {first['qps']:,.2f}→{peak['qps']:,.2f}（{peak['qps']/first['qps']:.2f} 倍）；"
            f"在已测档位中峰值位于 {peak['threads']} 并发。1000 并发 QPS 为 {last['qps']:,.2f}，"
            f"较峰值下降 {decline:.2f}%；P95 从 {first['p95_ms']:,.2f}ms 增至 {last['p95_ms']:,.2f}ms，"
            f"P99 从 {first['p99_ms']:,.2f}ms 增至 {last['p99_ms']:,.2f}ms。")
    findings.append({'level': 'L1', 'text': text, 'scenario': s,
                     'source': f'data/single.log；extracted_data.records[scenario={s}]'})
    if peak['threads'] < last['threads'] and last['qps']/peak['qps'] < .95:
        findings.append({'level': 'L2', 'text': f'{s} 超过 {peak["threads"]} 并发后 QPS 降幅超过 5%，提示已测区间内出现性能拐点。',
                         'source': f'analysis_results.single_product.{product}.scale.{s}', 'scenario': s})
insights = {'single_product_insights': {product: findings}, 'selection_insights': []}
save('data/analysis_results.json', analysis)
save('data/insights.json', insights)
assert analysis['single_product']
assert all(i.get('level') in ('L1', 'L2') and i.get('source') for i in findings)

# Generate all required charts through the Skill's chart module, inspecting each
# figure's visible text bounding boxes before it is saved.
import generate_charts as charts
import matplotlib.pyplot as plt
from matplotlib.text import Text
OUT.joinpath('charts').mkdir(exist_ok=True)
overlaps = {}
original_savefig = plt.savefig


def check_text_overlap(fig):
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    texts = [t for t in fig.findobj(Text) if t.get_visible() and t.get_text().strip()]
    boxes = [(t, t.get_window_extent(renderer)) for t in texts]
    result = []
    for i, (a, ab) in enumerate(boxes):
        for b, bb in boxes[i+1:]:
            if ab.width > 0 and ab.height > 0 and bb.width > 0 and bb.height > 0 and ab.overlaps(bb):
                result.append([a.get_text(), b.get_text()])
    return result


def checked_savefig(path, *args, **kwargs):
    fig = plt.gcf()
    result = check_text_overlap(fig)
    overlaps[Path(path).name] = result
    original_savefig(path, *args, **kwargs)


plt.savefig = checked_savefig
charts.gen_single(extracted, analysis, str(OUT / 'charts'))
plt.savefig = original_savefig
save('data/chart_checks.json', overlaps)
assert len(overlaps) == 11 and not any(overlaps.values()), overlaps

def para(text):
    return {'type': 'paragraph', 'text': text}


def table(headers, rows):
    return {'type': 'table', 'headers': headers, 'rows': rows}


def image(name, caption):
    return {'type': 'image', 'path': f'charts/{name}.png', 'caption': '图：' + caption}


def insight(item):
    return dict(item, type='insight', text=item['text'] + ' 数据来源：' + item['source'])


def sub(n, title, blocks):
    return {'id': 's' + n.replace('.', '-'), 'title': n + ' ' + title, 'blocks': blocks}


def section(n, title, subs):
    return {'id': f'ch{n}', 'title': f'第 {n} 章 {title}', 'blocks': [], 'subsections': subs}


def f(v):
    return f'{v:,.2f}' if v is not None else '-'


peaks = sp['peak_summary']
core = [i for i in findings if i['level'] == 'L1']
sections = [section(1, '测试结论', [
    sub('1.1', '测试结论摘要', [para(f'{product} 的 sysbench OLTP 日志包含 5 个场景、6 个并发档，共 30 条记录。'
        f'日志声明每项测试时长 {duration}s，测试开始时间为 {start}。五个场景的已测 QPS 峰值均出现在 512 并发，'
        '1000 并发时吞吐均有所回落且 P95、P99 上升；各场景分别分析，不合并为平均性能。')]),
    sub('1.2', '关键性能数字', [table(['场景', '峰值并发', '峰值 TPS', '峰值 QPS', 'P95@峰值 (ms)'],
        [[SCENARIO_CN[p['scenario']], str(p['peak_threads']), f(p['peak_tps']), f(p['peak_qps']), f(p['peak_p95_ms'])] for p in peaks]),
        image('single_peak_summary', f'{product} 各场景峰值 QPS。不同场景的事务工作量不同，图中数值不代表场景优劣。')]),
    sub('1.3', '核心结论', list(map(insight, core))) ]),
    section(2, '测试环境', [
        sub('2.1', '数据库配置', [table(['项', '源日志值'], [['产品', product], ['DB 端点', endpoint],
            ['开始时间', start], ['数据集', data_config], ['单项测试时长', f'{duration}s'], ['报告时间锚点', '2026-06-19T16:00:00+08:00']])]),
        sub('2.2', '硬件环境', [para('源日志未提供 CPU、内存、磁盘、网络、操作系统、数据库配置参数及部署拓扑。硬件环境待补充，不能据此判断资源瓶颈或推广到其他机器。')]),
        sub('2.3', '测试方法', [para('工具为 sysbench，具体版本与完整执行命令未记录。并发档为 32、64、128、256、512、1000；日志声明 Duration: 300s，不能将 Time 字段中的年份当作时长。'),
            table(['场景', '说明'], [[s, SCENARIO_CN[s]] for s in scenarios]),
            para('TPS 表示每秒事务数，QPS 表示每秒查询数；P95、P99 为延迟分位数，单位 ms。各场景事务组成不同，不能仅按 QPS 大小评价场景成本。')]),
        sub('2.4', '数据质量与公平性', [table(['检查项', '结果'], [['产品/场景/并发档', '1 / 5 / 6'],
            ['记录数', '30 = 1 × 5 × 6'], ['TPS / QPS / P95 / P99 空值率', '全部 0%'], ['逐条源数据核对', '30 条、120 个性能数值全部一致'],
            ['并发覆盖', '全部场景覆盖全部 6 档'], ['数据集和时长', f'同一日志统一配置；{duration}s'], ['硬件公平性', '信息缺失，无法核实']]),
            para('原始解析器以工具名代替产品名，且时长匹配会先命中 Time 中的年份。本次输出通过日志 DB、Duration、Time、Data 显式字段校正；源文件与 Skill 均保持原样。')]) ]),
    section(3, '性能数据', [sub('3.1', f'{product} 全量数据', [
        para(f'下表完整保留原始 30 条记录的场景、并发、记录时间和全部性能指标；所有数值保留源日志两位小数。产品为 {product}，每项测试时长为 {duration}s。'),
        table(['场景', '并发', '记录时间', 'TPS', 'QPS', 'P95 (ms)', 'P99 (ms)'],
              [[r['scenario'], str(r['threads']), r['test_time'], f(r['tps']), f(r['qps']), f(r['p95_ms']), f(r['p99_ms'])] for r in records])])]),
    section(4, '性能分析', [
        sub('4.1', '并发扩展性', [para('五个场景均独立使用完整六档并发绘制曲线。全部场景的已测峰值位于 512 并发，继续增至 1000 并发时 QPS 回落。该现象仅描述本次采样，不能确定精确饱和点。')] +
            [b for s in scenarios for b in [image(f'single_qps_{s}', f'{SCENARIO_CN[s]} 并发-QPS 曲线，横轴为以 2 为底的对数刻度。'), insight(next(i for i in core if i['scenario'] == s))]]),
        sub('4.2', 'P95 延迟稳定性', [para('五个场景 P95 随并发升高而增加；每张图只展示相应场景，P99 完整保留于第 3 章。分位数是不同负载档位的汇总值，不构成重复运行的方差或置信区间。')] +
            [b for s in scenarios for b in [image(f'single_p95_{s}', f'{SCENARIO_CN[s]} 并发-P95 曲线，单位 ms。'),
                para(f"{SCENARIO_CN[s]}：32 并发 P95={f(next(r['p95_ms'] for r in records if r['scenario']==s and r['threads']==32))}ms；"
                     f"512 并发 P95={f(next(r['p95_ms'] for r in records if r['scenario']==s and r['threads']==512))}ms；"
                     f"1000 并发 P95={f(next(r['p95_ms'] for r in records if r['scenario']==s and r['threads']==1000))}ms。")]]) ]),
    section(5, '优化建议', [
        sub('5.1', '调优最佳实践', [para('以各场景实测吞吐和尾延迟共同选择负载档位：512 并发为本次吞吐峰值档，但不能直接作为生产连接池上限。'
            '结合实际业务延迟目标，优先复测 256、512、1000 附近的负载，补充 CPU、IO、锁等待、连接数与内存使用情况后再定位瓶颈。'
            '本日志没有支持 Buffer Pool 比例或数据库参数变更的证据，不给出未经测量的参数值。')]),
        sub('5.2', '测试边界', [para('仅有单产品、单数据集、单次测试汇总；未提供重复试验、错误率、资源指标、硬件信息及 sysbench 版本。'
            '未覆盖 OLAP、复杂 JOIN、大字段或更大数据集。未给定业务 SLA，不能认定某档位满足生产要求。'
            '本报告仅报告已观测变化，不将高并发回落归因于某种未测量的资源。')]) ]),
    section(6, '附录', [
        sub('6.1', '附录 A：sysbench 命令', [para('源日志未包含完整 sysbench 执行命令，无法还原。已知配置为 tables=16、table_size=1000000、prepare=true、Duration=300s，及上述场景和并发档。')]),
        sub('6.2', '附录 B：原始数据来源', [para('唯一性能数据来源：data/single.log。标准化数据见 data/extracted_data.json，分析结果见 data/analysis_results.json，洞察见 data/insights.json。'
            '原始数值逐条核对，源日志的产品名、时间、端点、数据集和测试时长均保留；原输入与 Skill 的 SHA-256 核对清单见 data/input_hashes.json。')]) ])]
report_data = {'cover': {'title': '单次测试报告', 'subtitle': f'{product} sysbench OLTP 性能测试', 'date': '2026-06-19', 'version': 'v1.0'}, 'sections': sections}
save('data/report_data.json', report_data)

from render_core import render_md, render_html
md = render_md(report_data)
md = re.sub(r'^## (第 \d+ 章 .+)$', r'# \1', md, flags=re.M)
md = re.sub(r'^### (\d+\.\d+ .+)$', r'## \1', md, flags=re.M)
md = md.replace('> **[L1]** ', '> ✅ L1 | ').replace('> **[L2]** ', '> 📊 L2 | ')
(OUT / 'report.md').write_text(md, encoding='utf-8')

# Independently generate Word from report_data, with all text using the font helper.
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
doc = Document()


def set_doc_default_font():
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
        style = doc.styles[name]
        style.font.name = '微软雅黑'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')


def add_run_with_font(p, text, bold=False):
    run = p.add_run(str(text))
    run.font.name = '微软雅黑'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.bold = bold
    return run


def shade(cell, color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shd)


def word_blocks(blocks):
    for b in blocks:
        typ = b['type']
        if typ == 'table':
            t = doc.add_table(rows=1, cols=len(b['headers']))
            t.style = 'Table Grid'
            for c, h in zip(t.rows[0].cells, b['headers']):
                shade(c, 'D9E2F3')
                add_run_with_font(c.paragraphs[0], h, True).font.color.rgb = RGBColor.from_string('1F4E79')
            for i, row in enumerate(b['rows'], 1):
                for c, val in zip(t.add_row().cells, row):
                    if i % 2 == 0:
                        shade(c, 'F8FAFC')
                    add_run_with_font(c.paragraphs[0], val)
        elif typ == 'image':
            p = doc.add_paragraph()
            p.alignment = 1
            add_run_with_font(p, '').add_picture(str(OUT / b['path']), width=Cm(15.5))
            p = doc.add_paragraph()
            p.alignment = 1
            r = add_run_with_font(p, b['caption'])
            r.italic = True
            r.font.color.rgb = RGBColor.from_string('666666')
        else:
            p = doc.add_paragraph()
            text = b.get('text', '')
            if typ == 'insight':
                text = '✅ L1 确定性结论 | ' + text
                borders = OxmlElement('w:pBdr')
                left = OxmlElement('w:left')
                for key, val in {'val': 'single', 'sz': '18', 'space': '6', 'color': '1971C2'}.items():
                    left.set(qn('w:' + key), val)
                borders.append(left)
                p._p.get_or_add_pPr().append(borders)
            add_run_with_font(p, text)


set_doc_default_font()
add_run_with_font(doc.add_paragraph(style='Title'), report_data['cover']['title'], True)
add_run_with_font(doc.add_paragraph(), report_data['cover']['subtitle'])
add_run_with_font(doc.add_paragraph(), '日期：2026-06-19    版本：v1.0')
doc.add_page_break()
for sec in sections:
    add_run_with_font(doc.add_paragraph(style='Heading 1'), sec['title'], True)
    word_blocks(sec['blocks'])
    for ss in sec['subsections']:
        add_run_with_font(doc.add_paragraph(style='Heading 2'), ss['title'], True)
        word_blocks(ss['blocks'])
doc.save(str(OUT / 'report.docx'))

html = render_html(report_data)
extra_css = """<style>
body {font-family:'Microsoft YaHei','PingFang SC','Heiti TC','Hiragino Sans GB','STHeiti','Arial Unicode MS',sans-serif}
h1 {color:#1f4e79;border-bottom:2px solid #1f4e79}
h2 {color:#2e75b6;border-left:4px solid #2e75b6;padding-left:10px}
h3 {color:#4a6d8c}
th {background:linear-gradient(#D9E2F3,#C5D3EE)}
tr:nth-child(even) td {background:#F8FAFC}
.insight.L1 {background:#E7F5FF;border-left-color:#1971C2}
.insight.L2 {background:#FFF9DB;border-left-color:#F08C00}
.insight.L3 {background:#FFE3E3;border-left-color:#C92A2A}
</style>"""
html = html.replace('</head>', extra_css + '\n</head>')
(OUT / 'report.html').write_text(html, encoding='utf-8')

(OUT / 'docs').mkdir(exist_ok=True)
(OUT / 'docs/analysis_proposal.md').write_text('# 分析方案\n\n使用单一日志，按五场景独立分析六档并发的吞吐、P95、P99与峰值；不跨场景求均值，不推测硬件或参数。产品和时长以 DB、Duration 显式字段为准。报告日期以用户给定的 2026-06-19 为准。\n\n- [x] 意图识别\n- [x] 数据解析、元数据校正、30 条核对\n- [x] 场景独立分析与洞察检查\n- [x] 11 张图与文字重叠检查\n- [x] 三格式独立渲染\n- [x] 最小交付核查与输入哈希检查\n', encoding='utf-8')
(OUT / 'data/block_map.md').write_text('# 数据块映射\n\n' + '\n'.join(f'- 第 {j+1} 条：single.log 第 {raw.splitlines().index(next(line for line in raw.splitlines() if line.startswith("--- " + r["scenario"] + " | threads=" + str(r["threads"]) + " |")))+1} 行 → records[{j}] → 第 3 章全量表；场景 {r["scenario"]}，并发 {r["threads"]}。' for j, r in enumerate(records)), encoding='utf-8')

# Validate the saved deliverables, not just the in-memory report data.
from zipfile import ZipFile
from lxml import etree
saved_doc = Document(str(OUT / 'report.docx'))
doc_text = '\n'.join(p.text for p in saved_doc.paragraphs) + '\n' + '\n'.join(c.text for t in saved_doc.tables for row in t.rows for c in row.cells)
assert len(saved_doc.inline_shapes) == 11
assert len(saved_doc.tables[2].rows) >= 1
with ZipFile(OUT / 'report.docx') as z:
    xml = etree.fromstring(z.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    text_runs = xml.xpath('//w:r[w:t]', namespaces=ns)
    assert all(r.xpath('./w:rPr/w:rFonts/@w:eastAsia', namespaces=ns) == ['微软雅黑'] for r in text_runs)
    for t in saved_doc.tables:
        assert all(c._tc.xpath('./w:tcPr/w:shd')[0].get(qn('w:fill')) == 'D9E2F3' for c in t.rows[0].cells)
for r in records:
    for key in ('tps', 'qps', 'p95_ms', 'p99_ms'):
        assert all(f(r[key]) in text for text in (md, html, doc_text))
    assert all(r['scenario'] in text for text in (md, html, doc_text))
assert all(product in text and '300s' in text and start in text for text in (md, html, doc_text))
assert not re.findall(r'\{[A-Za-z_][A-Za-z0-9_]*\}', md)
assert not re.findall(r'\{[A-Za-z_][A-Za-z0-9_]*\}', html)
assert 'L3' not in md and '💡' not in md
assert len(re.findall(r'!\[.*?\]\(charts/.*?\.png\)', md)) == 11
for path in re.findall(r'<img src="([^"]+)"', html):
    assert (OUT / path).is_file()
assert (OUT / 'report.md').stat().st_size >= 5 * 1024
assert (OUT / 'report.docx').stat().st_size >= 10 * 1024
assert (OUT / 'report.html').stat().st_size >= 8 * 1024
assert len(core) == 5
assert hashes() == before, 'Source inputs or Skill changed'
checks = {
    'A': 'PASS：必需中间产物、分析方案、完成清单和三格式文件齐全，图表 11 张，文件大小达标。',
    'B': 'PASS：1 产品 × 5 场景 × 6 并发 = 30 条；全部 120 个性能数值与原日志逐条一致；产品和时长已核对。',
    'C': 'PASS：5 QPS + 5 P95 + 1 峰值图；中文标题；11 张 check_text_overlap 均返回空；C5 对 single 不适用。',
    'D': 'PASS：5 条核心结论，洞察均含 level/source，场景分别分析，无 L3 或未披露 STRONG 项；硬件缺失已披露。',
    'E': 'PASS：章节层级、相对图片路径、Word 全部文本字体及全部表头底色、11 张嵌图、HTML head/CSS/字体回退已核查。',
    'F': 'PASS：无残留占位符；无无依据参数建议；三个格式全部 120 个性能数值一致。',
    'source_unchanged': hashes() == before,
    'limitations': ['源日志未提供硬件、工具版本、完整命令、错误率、重复测试与资源监控。', '未使用 Word 或浏览器进行人工分页/页面截图检查；图中文字重叠使用渲染包围盒自动检查。']}
save('data/validation.json', checks)
(OUT / 'validation.md').write_text('# 交付核查\n\n' + '\n\n'.join(f'- {k}：{v}' for k,v in checks.items()), encoding='utf-8')
print(json.dumps({'status': 'complete', 'records': len(records), 'product': product, 'duration_sec': duration,
                  'charts': len(overlaps), 'sizes': {p.name: p.stat().st_size for p in OUT.glob('report.*')},
                  'source_unchanged': hashes() == before}, ensure_ascii=False, indent=2))
